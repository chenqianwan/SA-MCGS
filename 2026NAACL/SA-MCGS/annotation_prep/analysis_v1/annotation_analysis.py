#!/usr/bin/env python3
"""Offline expert-annotation analysis for the SA-MCGS blind audit.

The pipeline is deliberately fail-closed.  It never calls a model, never derives
legacy labels, and never chooses an adjudicated label from system outcomes.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import itertools
import json
import math
import random
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence


EXPORT_SCHEMA = "blind_annotation.export.v2"
BINARY_PROTOCOL = "binary_judgment_v1"
MERGED_SCHEMA = "annotation_analysis.merged.v1"
IAA_SCHEMA = "annotation_analysis.iaa.v1"
ADJUDICATION_SCHEMA = "annotation_analysis.adjudication.v1"
FROZEN_SCHEMA = "annotation_analysis.frozen_labels.v1"
LEGACY_SCHEMA = "annotation_analysis.legacy_labels.v1"
LEGACY_SOURCE_SCHEMA = "annotation_analysis.legacy_source_records.v1"
SYSTEM_OUTPUT_SCHEMA = "annotation_analysis.system_outputs.v1"
CLEANLINESS_SCHEMA = "annotation_analysis.cleanliness.v1"
CLUSTER_SUMMARY_SCHEMA = "annotation_analysis.clustered_summary.v1"
LAYOUT_SCHEMA = "annotation_analysis.layout_audit.v1"
TOOL_VERSION = "sa_mcgs_annotation_analysis.v2"
SHA256_PREFIXED_RE = re.compile(r"^sha256:[0-9a-f]{64}$")
SHA256_BARE_RE = re.compile(r"^[0-9a-f]{64}$")

ADJUDICATION_SELECTION_RULE = [
    "no 2-of-3 majority on a categorical field",
    "any target=insufficient rating",
    "any confidence <= 2",
    "positive/sufficient provisional label without a 2-of-3 minimal-evidence node",
    "positive/sufficient provisional label without a 2-of-3 repair-packet node",
]

EXPERT_IDS = ("expert_1", "expert_2", "expert_3")
EXPERT_LABEL_TO_ID = {
    "专家 1": "expert_1",
    "专家 2": "expert_2",
    "专家 3": "expert_3",
    **{expert_id: expert_id for expert_id in EXPERT_IDS},
}
TARGET_LABELS = ("present", "absent", "insufficient")
SUFFICIENCY_LABELS = ("yes", "partial", "no", "not_applicable")
CONTEXT_LABELS = ("yes", "no", "uncertain")
SECTION_ROLES = (
    "conflict_evidence",
    "supporting_context",
    "ordinary_context",
    "uncertain",
)
PREDICTIONS = ("present", "absent", "insufficient", "abstain", "invalid")
FORMAL_STRATUM_COUNTS = {"pre18": 18, "post_new24": 24, "post_legacy12": 12}

CASE_XLSX_REQUIRED = {
    "专家编号",
    "匿名标注者编号",
    "展示顺序",
    "样本编号",
    "目标冲突代码",
    "材料充分性代码",
    "更多上下文代码",
    "判断信心_1到5",
    "完成状态",
    "manifest_hash",
    "translation_pack_hash",
}
SECTION_XLSX_REQUIRED = {
    "专家编号",
    "匿名标注者编号",
    "样本编号",
    "段落编号",
    "段落作用代码",
    "是否最小证据",
    "是否修复材料",
}
METADATA_XLSX_REQUIRED = {
    "研究编号",
    "专家编号",
    "匿名标注者编号",
    "manifest_hash",
    "translation_pack_hash",
}


class ValidationError(ValueError):
    """Raised when an artifact fails an analysis gate."""


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationError(f"Cannot read JSON {path}: {exc}") from exc


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def canonical_sha256(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return "sha256:" + hashlib.sha256(encoded).hexdigest()


def normalize_text(value: Any) -> str:
    """Match the canonical post-material generator's whitespace normalization."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def canonical_post_content_fingerprint(record: dict[str, Any]) -> str:
    """Recompute the generator-defined fingerprint from ordered canonical node text."""
    case_uid = string(record.get("post_case_uid"))
    require(case_uid, "Post record has no post_case_uid")
    nodes = record.get("nodes")
    require(isinstance(nodes, list) and nodes, f"Post {case_uid} has no canonical nodes")
    canonical_nodes = []
    seen: set[str] = set()
    for index, node in enumerate(nodes):
        require(isinstance(node, dict), f"Post {case_uid} node {index} is not an object")
        node_id = string(node.get("id"))
        require(node_id and node_id not in seen, f"Post {case_uid} has empty/duplicate node ID")
        require("content" in node, f"Post {case_uid}/{node_id} has no canonical content field")
        seen.add(node_id)
        canonical_nodes.append({
            "id": node_id,
            "normalized_text": normalize_text(node.get("content")),
        })
    return canonical_sha256({
        "version": "post_injected_text_content_v1",
        "post_case_uid": case_uid,
        "nodes": canonical_nodes,
    })


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def seal(value: dict[str, Any], field: str = "artifact_sha256") -> dict[str, Any]:
    result = copy.deepcopy(value)
    result.pop(field, None)
    result[field] = canonical_sha256(result)
    return result


def verify_seal(value: dict[str, Any], field: str = "artifact_sha256") -> None:
    expected = value.get(field)
    if not isinstance(expected, str):
        raise ValidationError(f"Artifact has no {field}")
    unsigned = copy.deepcopy(value)
    unsigned.pop(field, None)
    actual = canonical_sha256(unsigned)
    if actual != expected:
        raise ValidationError(f"Artifact hash mismatch: expected {expected}, got {actual}")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def string(value: Any) -> str:
    return "" if value is None else str(value).strip()


def unique_strings(value: Any, label: str) -> list[str]:
    require(isinstance(value, list), f"{label} must be a list")
    result = [string(item) for item in value]
    require(all(result), f"{label} contains an empty ID")
    require(len(result) == len(set(result)), f"{label} contains duplicate IDs")
    return result


def is_synthetic(payload: dict[str, Any]) -> bool:
    return payload.get("synthetic_fixture") is True


def validate_iso_datetime(value: str, label: str) -> str:
    require(bool(value), f"{label} is required")
    try:
        datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValidationError(f"{label} must be ISO-8601: {value!r}") from exc
    return value


def require_sha256(value: Any, label: str, *, prefixed: bool = True) -> str:
    digest = string(value)
    pattern = SHA256_PREFIXED_RE if prefixed else SHA256_BARE_RE
    require(bool(pattern.fullmatch(digest)), f"{label} must be a lowercase SHA-256 digest")
    return digest


def recompute_manifest_hash(manifest: dict[str, Any]) -> str:
    unsigned = copy.deepcopy(manifest)
    unsigned["manifest_hash"] = ""
    encoded = json.dumps(
        unsigned,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def source_text_hash(title: str, text: str) -> str:
    return canonical_sha256({"title": title, "text": text})


def load_manifest(path: Path) -> dict[str, Any]:
    manifest = read_json(path)
    require(isinstance(manifest, dict), "Public manifest must be an object")
    require(manifest.get("schema_version") == "blind_annotation.v1", "Wrong manifest schema")
    require(bool(string(manifest.get("study_id"))), "Manifest has no study_id")
    stored_manifest_hash = require_sha256(manifest.get("manifest_hash"), "manifest_hash", prefixed=False)
    actual_manifest_hash = recompute_manifest_hash(manifest)
    require(
        stored_manifest_hash == actual_manifest_hash,
        f"Manifest body hash mismatch: stored {stored_manifest_hash}, recomputed {actual_manifest_hash}",
    )
    cases = manifest.get("cases")
    require(isinstance(cases, list) and cases, "Manifest cases must be a non-empty list")
    case_ids: list[str] = []
    for item in cases:
        require(isinstance(item, dict), "Manifest case is not an object")
        case_id = string(item.get("presentation_id"))
        require(case_id and case_id not in case_ids, f"Invalid/duplicate presentation_id {case_id!r}")
        case_ids.append(case_id)
        sections = item.get("sections")
        require(isinstance(sections, list) and sections, f"{case_id} has no sections")
        require(
            all(
                isinstance(section, dict)
                and isinstance(section.get("title"), str)
                and isinstance(section.get("text"), str)
                for section in sections
            ),
            f"{case_id} section title/text fields must be strings",
        )
        section_ids = [string(section.get("section_id")) for section in sections if isinstance(section, dict)]
        require(len(section_ids) == len(sections), f"{case_id} has an invalid section")
        require(all(section_ids) and len(section_ids) == len(set(section_ids)), f"{case_id} section IDs invalid")
    orders = manifest.get("expert_orders")
    require(isinstance(orders, dict) and set(orders) == set(EXPERT_IDS), "Manifest expert_orders invalid")
    for expert_id in EXPERT_IDS:
        order = unique_strings(orders[expert_id], f"expert_orders.{expert_id}")
        require(set(order) == set(case_ids), f"Frozen order for {expert_id} is not a case permutation")
    return manifest


def load_translation_pack(path: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    translations = read_json(path)
    require(isinstance(translations, dict), "Translation pack must be an object")
    require(
        translations.get("schema_version") == "annotation_translation.zh-CN.v1",
        "Wrong translation schema",
    )
    require(
        translations.get("source_manifest_hash") == manifest["manifest_hash"],
        "Translation pack source_manifest_hash mismatch",
    )
    stored_hash = require_sha256(
        translations.get("translation_pack_hash"), "translation_pack_hash"
    )
    unsigned = copy.deepcopy(translations)
    unsigned.pop("translation_pack_hash", None)
    actual_hash = canonical_sha256(unsigned)
    require(
        stored_hash == actual_hash,
        f"Translation pack body hash mismatch: stored {stored_hash}, recomputed {actual_hash}",
    )
    translated_cases = translations.get("cases")
    require(isinstance(translated_cases, dict), "Translation pack cases must be an object")
    manifest_cases = {case["presentation_id"]: case for case in manifest["cases"]}
    require(
        set(translated_cases) == set(manifest_cases),
        "Translation pack cases are not a manifest bijection",
    )
    for case_id, manifest_case in manifest_cases.items():
        translated_case = translated_cases[case_id]
        require(isinstance(translated_case, dict), f"Translation case {case_id} invalid")
        translated_sections = translated_case.get("sections")
        require(isinstance(translated_sections, dict), f"Translation case {case_id} has no sections")
        manifest_sections = {section["section_id"]: section for section in manifest_case["sections"]}
        require(
            set(translated_sections) == set(manifest_sections),
            f"Translation sections for {case_id} are not a manifest bijection",
        )
        for section_id, source in manifest_sections.items():
            translated = translated_sections[section_id]
            require(isinstance(translated, dict), f"Translation {case_id}/{section_id} invalid")
            expected_source_hash = source_text_hash(
                source["title"], source["text"]
            )
            require(
                translated.get("source_hash") == expected_source_hash,
                f"Translation source hash mismatch for {case_id}/{section_id}",
            )
            require(
                isinstance(translated.get("title_zh"), str)
                and isinstance(translated.get("text_zh"), str),
                f"Translation text fields invalid for {case_id}/{section_id}",
            )
    return translations


def current_stratum(master_record: dict[str, Any]) -> str:
    action = string(master_record.get("annotation_action"))
    require(action == "new_annotation", f"Current record action must be new_annotation, got {action!r}")
    if master_record.get("unmodified") is True:
        return "pre18"
    require(bool(string(master_record.get("post_case_uid"))), "New post record has no post_case_uid")
    return "post_new24"


def _node_ids(nodes: Any, label: str) -> list[str]:
    require(isinstance(nodes, list) and nodes, f"{label} must be a non-empty list")
    result = []
    for node in nodes:
        node_id = string(node.get("id")) if isinstance(node, dict) else string(node)
        require(node_id, f"{label} has a node without ID")
        result.append(node_id)
    require(len(result) == len(set(result)), f"{label} has duplicate node IDs")
    return result


def _resolve_linked_path(raw: Any, owner_path: Path, label: str) -> Path:
    value = string(raw)
    require(value, f"{label} path is required")
    path = Path(value)
    if not path.is_absolute():
        path = owner_path.parent / path
    path = path.resolve()
    require(path.is_file(), f"{label} file does not exist: {path}")
    return path


def _validate_codebook_source_master(
    codebook: dict[str, Any], codebook_path: Path
) -> tuple[Path, dict[str, Any]]:
    source_path = _resolve_linked_path(
        codebook.get("source_master"), codebook_path, "Codebook source_master"
    )
    source = read_json(source_path)
    require(
        isinstance(source, dict) and source.get("schema_version") == "annotation_master.v1",
        "Codebook source_master schema invalid",
    )
    source_records = source.get("cases")
    require(isinstance(source_records, list), "Codebook source_master cases must be a list")
    current_records = codebook.get("records")
    reused_records = codebook.get("reused_records")
    skipped_records = codebook.get("skipped_records")
    require(isinstance(current_records, list), "Codebook records must be a list")
    require(isinstance(reused_records, list), "Codebook reused_records must be a list")
    require(isinstance(skipped_records, list), "Codebook skipped_records must be a list")
    source_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    for row in source_records:
        require(isinstance(row, dict), "source_master case is not an object")
        if string(row.get("pre_source_id")):
            key = ("pre", string(row.get("pre_source_id")))
        else:
            require(string(row.get("post_case_uid")), "source_master case has no stable identity")
            key = ("post", string(row.get("post_case_uid")))
        require(key not in source_by_key, f"source_master duplicate stable identity {key}")
        source_by_key[key] = row
    materialized = []
    for row in current_records:
        require(isinstance(row, dict) and isinstance(row.get("master_record"), dict), "Codebook current master_record invalid")
        master = row["master_record"]
        if string(master.get("pre_source_id")):
            key = ("pre", string(master.get("pre_source_id")))
            expected_original = string(master.get("pre_source_id"))
        else:
            key = ("post", string(master.get("post_case_uid")))
            expected_original = string(master.get("canonical_case_id"))
        require(key in source_by_key, f"Codebook wrapper references unknown source_master identity {key}")
        require(
            canonical_sha256(master) == canonical_sha256(source_by_key[key]),
            f"Codebook wrapper/master binding differs from source_master for {key}",
        )
        require(
            string(row.get("original_case_id")) == expected_original,
            f"Codebook wrapper original_case_id does not bind its embedded master for {key}",
        )
        materialized.append(master)
    materialized.extend(reused_records)
    materialized.extend(skipped_records)
    source_digests = Counter(canonical_sha256(row) for row in source_records)
    materialized_digests = Counter(canonical_sha256(row) for row in materialized)
    require(
        source_digests == materialized_digests,
        "Codebook records are not an exact multiset copy of source_master cases",
    )
    return source_path, source


def _load_post_identity_catalog(
    source_master_path: Path,
    source_master: dict[str, Any],
    synthetic: bool,
) -> tuple[Path, dict[str, dict[str, Any]]]:
    """Bind post endpoints/bridges to the outcome-blind canonical selection record."""
    if synthetic:
        catalog_path = source_master_path.with_name("synthetic_post_identity_catalog.json")
    else:
        catalog_path = source_master_path.parent.parent / "candidates" / "master_manifest.json"
    catalog_path = catalog_path.resolve()
    require(catalog_path.is_file(), f"Canonical post identity catalog is missing: {catalog_path}")
    catalog = read_json(catalog_path)
    require(isinstance(catalog, dict), "Canonical post identity catalog must be an object")
    require(is_synthetic(catalog) == synthetic, "Post identity catalog synthetic marker mismatch")
    if synthetic:
        require(
            catalog.get("schema_version") == "annotation_analysis.synthetic_post_identity_catalog.v1",
            "Synthetic post identity catalog schema invalid",
        )
    else:
        require(
            bool(string(catalog.get("manifest_version"))) and catalog.get("internal_only") is True,
            "Formal post identity catalog schema invalid",
        )
    raw_rows = catalog.get("post_presentations")
    require(isinstance(raw_rows, list), "Post identity catalog post_presentations must be a list")
    catalog_by_uid: dict[str, dict[str, Any]] = {}
    for raw in raw_rows:
        require(isinstance(raw, dict), "Post identity catalog row is not an object")
        uid = string(raw.get("post_case_uid"))
        require(uid and uid not in catalog_by_uid, "Post identity catalog case UIDs empty/duplicate")
        catalog_by_uid[uid] = raw

    post_master = {
        string(row.get("post_case_uid")): row
        for row in source_master.get("cases", [])
        if isinstance(row, dict) and string(row.get("post_case_uid"))
    }
    require(set(catalog_by_uid) == set(post_master), "Post identity catalog does not exactly cover source_master post cases")
    result: dict[str, dict[str, Any]] = {}
    for uid, master in post_master.items():
        row = catalog_by_uid[uid]
        for field in ("canonical_case_id", "source_physical_fingerprint", "domain", "template"):
            require(
                string(row.get(field)) == string(master.get(field)),
                f"Post identity catalog {uid} field {field} differs from source_master",
            )
        node_ids = set(_node_ids(master.get("nodes"), f"post {uid}.nodes"))
        injection = row.get("injection_identity")
        require(isinstance(injection, dict), f"Post identity catalog {uid} lacks injection_identity")
        root = string(injection.get("root_node"))
        witness = string(injection.get("witness_node"))
        require(root and witness, f"Post identity catalog {uid} lacks endpoint identities")
        endpoints = unique_strings(injection.get("risk_nodes"), f"post catalog/{uid}.risk_nodes")
        require(set(endpoints) == {root, witness}, f"Post identity catalog {uid} risk_nodes differ from root/witness")
        bridge = string(injection.get("bridge_node"))
        bridge_nodes = [bridge] if bridge else []
        evidence_nodes = unique_strings(injection.get("evidence_nodes"), f"post catalog/{uid}.evidence_nodes")
        require(
            set(evidence_nodes) == set(endpoints) | set(bridge_nodes),
            f"Post identity catalog {uid} evidence_nodes differ from endpoints plus bridge",
        )
        require(
            set(endpoints) | set(bridge_nodes) <= node_ids,
            f"Post identity catalog {uid} endpoint/bridge outside canonical node universe",
        )
        result[uid] = {
            "endpoint_node_ids": sorted(endpoints),
            "bridge_node_ids": sorted(bridge_nodes),
            "post_identity_record_sha256": canonical_sha256(row),
        }
    return catalog_path, result


def load_layout(
    manifest_path: Path,
    codebook_path: Path,
    translations_path: Path,
) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    translations = load_translation_pack(translations_path, manifest)
    codebook = read_json(codebook_path)
    require(isinstance(codebook, dict), "Private codebook must be an object")
    require(codebook.get("schema_version") == "blind_annotation.codebook.v1", "Wrong codebook schema")
    require(codebook.get("study_id") == manifest["study_id"], "Manifest/codebook study_id mismatch")
    require(codebook.get("manifest_hash") == manifest["manifest_hash"], "Manifest/codebook hash mismatch")
    synthetic = is_synthetic(codebook) or is_synthetic(manifest) or is_synthetic(translations)
    require(
        is_synthetic(codebook) == is_synthetic(manifest) == is_synthetic(translations),
        "Synthetic marker differs between manifest/codebook/translation pack",
    )
    source_master_path, source_master = _validate_codebook_source_master(codebook, codebook_path)
    post_catalog_path, post_identity = _load_post_identity_catalog(
        source_master_path, source_master, synthetic
    )

    public_cases = {item["presentation_id"]: item for item in manifest["cases"]}
    records = codebook.get("records")
    require(isinstance(records, list), "Codebook records must be a list")
    require(len(records) == len(public_cases), "Manifest/codebook current-case count mismatch")
    current: dict[str, dict[str, Any]] = {}
    for record in records:
        require(isinstance(record, dict), "Codebook record is not an object")
        presentation_id = string(record.get("presentation_id"))
        require(presentation_id in public_cases, f"Unknown codebook presentation {presentation_id!r}")
        require(presentation_id not in current, f"Duplicate codebook presentation {presentation_id}")
        master = record.get("master_record")
        require(isinstance(master, dict), f"{presentation_id} has no master_record")
        stratum = current_stratum(master)
        source_scc_uid = string(master.get("source_physical_fingerprint"))
        require(source_scc_uid, f"{presentation_id} has no source physical SCC identity")
        if not synthetic:
            require_sha256(source_scc_uid, f"{presentation_id}.source_physical_fingerprint")
        section_map = record.get("section_map")
        require(isinstance(section_map, list), f"{presentation_id} has no section_map")
        mapping: dict[str, str] = {}
        for item in section_map:
            require(isinstance(item, dict), f"{presentation_id} section_map item invalid")
            section_id = string(item.get("section_id"))
            node_id = string(item.get("original_node_id"))
            require(section_id and node_id and section_id not in mapping, f"{presentation_id} section map invalid")
            mapping[section_id] = node_id
        require(
            len(set(mapping.values())) == len(mapping),
            f"{presentation_id} section-to-node map is not one-to-one",
        )
        expected_sections = {section["section_id"] for section in public_cases[presentation_id]["sections"]}
        require(set(mapping) == expected_sections, f"{presentation_id} section map does not match public manifest")
        original_case_id = string(record.get("original_case_id"))
        require(original_case_id, f"{presentation_id} has no original_case_id")
        canonical_nodes = master.get("clauses") if stratum == "pre18" else master.get("nodes")
        canonical_node_ids = _node_ids(canonical_nodes, f"{presentation_id}.canonical_nodes")
        require(
            set(mapping.values()) == set(canonical_node_ids),
            f"{presentation_id} section map does not bijectively cover canonical nodes",
        )
        if stratum == "post_new24":
            post_content_fingerprint = string(master.get("post_content_fingerprint"))
            require_sha256(post_content_fingerprint, f"{presentation_id}.post_content_fingerprint")
            recomputed_post_fingerprint = canonical_post_content_fingerprint(master)
            require(
                post_content_fingerprint == recomputed_post_fingerprint,
                f"{presentation_id} post_content_fingerprint differs from canonical node text",
            )
            post_meta = post_identity[string(master.get("post_case_uid"))]
        else:
            post_content_fingerprint = None
            post_meta = {"endpoint_node_ids": [], "bridge_node_ids": [], "post_identity_record_sha256": None}
        current[presentation_id] = {
            "case_id": presentation_id,
            "original_case_id": original_case_id,
            "case_uid": string(master.get("post_case_uid") or master.get("pre_source_id") or original_case_id),
            "stratum": stratum,
            "domain": string(master.get("domain")),
            "source_scc_uid": source_scc_uid,
            "post_content_fingerprint": post_content_fingerprint,
            "section_id_to_node_id": mapping,
            "node_universe": sorted(mapping.values()),
            **post_meta,
        }

    reused = codebook.get("reused_records")
    require(isinstance(reused, list), "Codebook reused_records must be a list")
    legacy: dict[str, dict[str, Any]] = {}
    for record in reused:
        require(isinstance(record, dict), "Legacy codebook record is not an object")
        require(record.get("annotation_action") == "reuse_legacy_label", "Legacy record action mismatch")
        case_id = string(record.get("post_case_uid"))
        require(case_id and case_id not in legacy, f"Invalid/duplicate legacy post_case_uid {case_id!r}")
        source_scc_uid = string(record.get("source_physical_fingerprint"))
        require(source_scc_uid, f"Legacy case {case_id} has no source SCC identity")
        post_content_fingerprint = string(record.get("post_content_fingerprint"))
        require(post_content_fingerprint, f"Legacy case {case_id} has no post_content_fingerprint")
        legacy_case_id = string(record.get("legacy_case_id"))
        canonical_case_id = string(record.get("canonical_case_id"))
        require(legacy_case_id and canonical_case_id, f"Legacy case {case_id} has incomplete case identity")
        require(record.get("legacy_material_text_match") is True, f"Legacy case {case_id} is not certified text-identical")
        if not synthetic:
            require_sha256(source_scc_uid, f"legacy/{case_id}.source_physical_fingerprint")
        require_sha256(post_content_fingerprint, f"legacy/{case_id}.post_content_fingerprint")
        recomputed_post_fingerprint = canonical_post_content_fingerprint(record)
        require(
            post_content_fingerprint == recomputed_post_fingerprint,
            f"Legacy case {case_id} post_content_fingerprint differs from canonical node text",
        )
        post_meta = post_identity[case_id]
        legacy[case_id] = {
            "case_id": case_id,
            "original_case_id": canonical_case_id,
            "legacy_case_id": legacy_case_id,
            "case_uid": case_id,
            "stratum": "post_legacy12",
            "domain": string(record.get("domain")),
            "source_scc_uid": source_scc_uid,
            "post_content_fingerprint": post_content_fingerprint,
            "node_universe": sorted(_node_ids(record.get("nodes"), f"legacy {case_id}.nodes")),
            **post_meta,
        }

    counts = Counter(item["stratum"] for item in current.values())
    counts["post_legacy12"] = len(legacy)
    if synthetic:
        require(all(counts[key] > 0 for key in FORMAL_STRATUM_COUNTS), "Synthetic fixture must exercise all three strata")
    else:
        require(dict(counts) == FORMAL_STRATUM_COUNTS, f"Formal strata must be {FORMAL_STRATUM_COUNTS}, got {dict(counts)}")
    require(set(current) == set(public_cases), "Codebook is not a bijection over public cases")
    require(not (set(current) & set(legacy)), "Current and legacy primary case IDs must be globally disjoint")
    alias_owner: dict[str, str] = {}
    for identity in list(current.values()) + list(legacy.values()):
        owner = identity["case_id"]
        aliases = {
            owner,
            string(identity.get("original_case_id")),
            string(identity.get("case_uid")),
            string(identity.get("legacy_case_id")),
        } - {""}
        for alias in aliases:
            require(
                alias not in alias_owner or alias_owner[alias] == owner,
                f"Global case alias collision {alias!r}: {alias_owner.get(alias)!r} vs {owner!r}",
            )
            alias_owner[alias] = owner
    source_identity = {
        "manifest_file_sha256": file_sha256(manifest_path),
        "manifest_body_hash": manifest["manifest_hash"],
        "codebook_file_sha256": file_sha256(codebook_path),
        "source_master_path": str(source_master_path),
        "source_master_file_sha256": file_sha256(source_master_path),
        "post_identity_catalog_path": str(post_catalog_path),
        "post_identity_catalog_file_sha256": file_sha256(post_catalog_path),
        "translations_file_sha256": file_sha256(translations_path),
        "translation_pack_hash": translations["translation_pack_hash"],
        "identity_map_sha256": canonical_sha256({
            "current": current,
            "legacy": legacy,
        }),
    }
    return {
        "manifest": manifest,
        "codebook": codebook,
        "translations": translations,
        "synthetic_fixture": synthetic,
        "current": current,
        "legacy": legacy,
        "counts": dict(counts),
        "source_identity": source_identity,
    }


def layout_audit(
    manifest_path: Path,
    codebook_path: Path,
    translations_path: Path,
) -> dict[str, Any]:
    layout = load_layout(manifest_path, codebook_path, translations_path)
    current = list(layout["current"].values())
    legacy = list(layout["legacy"].values())
    payload = {
        "schema_version": LAYOUT_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": layout["synthetic_fixture"],
        "inputs": {
            "manifest": str(manifest_path.resolve()),
            "manifest_file_sha256": file_sha256(manifest_path),
            "codebook": str(codebook_path.resolve()),
            "codebook_file_sha256": file_sha256(codebook_path),
            "translations": str(translations_path.resolve()),
            "translations_file_sha256": file_sha256(translations_path),
            "source_master": layout["source_identity"]["source_master_path"],
            "source_master_file_sha256": layout["source_identity"]["source_master_file_sha256"],
            "post_identity_catalog": layout["source_identity"]["post_identity_catalog_path"],
            "post_identity_catalog_file_sha256": layout["source_identity"]["post_identity_catalog_file_sha256"],
        },
        "study_id": layout["manifest"]["study_id"],
        "manifest_hash": layout["manifest"]["manifest_hash"],
        "translation_pack_hash": layout["translations"]["translation_pack_hash"],
        "source_identity": layout["source_identity"],
        "stratum_policy": {
            "pre18": "current blind export; unmodified source SCC",
            "post_new24": "current blind export; newly annotated injected variant",
            "post_legacy12": "separate normalized legacy-label input; never inferred or pooled",
        },
        "counts": layout["counts"],
        "unique_source_scc": {
            key: len({row["source_scc_uid"] for row in current + legacy if row["stratum"] == key})
            for key in FORMAL_STRATUM_COUNTS
        },
        "public_sections": sum(len(item["sections"]) for item in layout["manifest"]["cases"]),
        "export_contract": {
            "json_schema": EXPORT_SCHEMA,
            "xlsx_sheets": ["样本整体标注", "逐段标注", "版本信息"],
            "xlsx_required_case_headers": sorted(CASE_XLSX_REQUIRED),
            "xlsx_required_section_headers": sorted(SECTION_XLSX_REQUIRED),
            "xlsx_required_metadata_headers": sorted(METADATA_XLSX_REQUIRED),
        },
        "gates": {
            "three_distinct_expert_slots": True,
            "manifest_and_translation_hash_equality": True,
            "exact_case_and_section_bijection": True,
            "wrapper_to_master_exact_case_binding": True,
            "canonical_post_text_fingerprint_recomputed": True,
            "canonical_endpoint_and_bridge_identity_bound": True,
            "formal_cardinality": FORMAL_STRATUM_COUNTS,
            "legacy_labels_must_be_explicit": True,
            "formal_legacy12_freeze_gate": "BLOCKED_UNTIL_EXPLICIT_VERIFIED_LEGACY_LABEL_ARTIFACT",
            "source_scc_is_cluster_unit": True,
        },
    }
    return seal(payload)


def _sheet_records(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    while rows and not any(cell not in (None, "") for cell in rows[-1]):
        rows.pop()
    require(bool(rows), f"Worksheet {sheet.title!r} is empty")
    headers = [string(cell) for cell in rows[0]]
    require(all(headers) and len(headers) == len(set(headers)), f"Worksheet {sheet.title!r} has invalid headers")
    records = []
    for raw in rows[1:]:
        if not any(cell not in (None, "") for cell in raw):
            continue
        padded = list(raw) + [None] * (len(headers) - len(raw))
        records.append(dict(zip(headers, padded[: len(headers)])))
    return records


def _require_headers(rows: list[dict[str, Any]], expected: set[str], label: str) -> None:
    require(bool(rows), f"{label} has no data rows")
    missing = expected - set(rows[0])
    require(not missing, f"{label} is missing headers: {sorted(missing)}")


def _expert_id(value: Any) -> str:
    raw = string(value)
    require(raw in EXPERT_LABEL_TO_ID, f"Unknown expert label {raw!r}")
    return EXPERT_LABEL_TO_ID[raw]


def xlsx_tables_to_payload(
    case_rows: list[dict[str, Any]],
    section_rows: list[dict[str, Any]],
    metadata_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    """Convert the current static page's three worksheet tables to export.v2."""
    for row in [*metadata_rows, *case_rows, *section_rows]:
        require(
            row.get("annotation_mode") in (None, "", "blind_human_audit")
            and not row.get("hint_pack_hash")
            and "辅助" not in string(row.get("标注模式")),
            "Assisted XLSX cannot be converted into a formal blind export",
        )
        require(
            not row.get("annotation_protocol"),
            "Explicit annotation protocol requires its own analysis; legacy XLSX merge cannot reinterpret it",
        )
    _require_headers(case_rows, CASE_XLSX_REQUIRED, "样本整体标注")
    _require_headers(section_rows, SECTION_XLSX_REQUIRED, "逐段标注")
    _require_headers(metadata_rows, METADATA_XLSX_REQUIRED, "版本信息")
    require(len(metadata_rows) == 1, "版本信息 must contain exactly one row")
    metadata = metadata_rows[0]
    expert_id = _expert_id(metadata["专家编号"])
    reviewer_id = string(metadata["匿名标注者编号"])
    require(reviewer_id, "XLSX has no anonymous reviewer ID")
    require(all(_expert_id(row["专家编号"]) == expert_id for row in case_rows), "Mixed experts in case sheet")
    require(all(_expert_id(row["专家编号"]) == expert_id for row in section_rows), "Mixed experts in section sheet")

    by_case_sections: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in section_rows:
        require(
            string(row["匿名标注者编号"]) == reviewer_id,
            "Reviewer ID differs inside XLSX section sheet",
        )
        case_id = string(row["样本编号"])
        section_id = string(row["段落编号"])
        require(case_id and section_id and section_id not in by_case_sections[case_id], "Duplicate/empty XLSX section row")
        minimal = string(row["是否最小证据"])
        repair = string(row["是否修复材料"])
        require(minimal in {"是", "否"} and repair in {"是", "否"}, "XLSX boolean cells must be 是/否")
        by_case_sections[case_id][section_id] = {
            "role": string(row["段落作用代码"]),
            "role_label": string(row.get("段落作用")),
            "minimal_evidence": minimal == "是",
            "repair_packet": repair == "是",
            "note": string(row.get("段落备注")),
        }

    ordered_rows = sorted(case_rows, key=lambda row: int(row["展示顺序"]))
    order: list[str] = []
    annotations: dict[str, dict[str, Any]] = {}
    for expected_index, row in enumerate(ordered_rows, start=1):
        require(int(row["展示顺序"]) == expected_index, "XLSX display order must be contiguous from 1")
        case_id = string(row["样本编号"])
        require(case_id and case_id not in annotations, "Duplicate/empty XLSX case row")
        require(string(row["完成状态"]) == "已完成", f"XLSX case {case_id} is incomplete")
        require(string(row["匿名标注者编号"]) == reviewer_id, "Reviewer ID differs inside XLSX")
        order.append(case_id)
        annotations[case_id] = {
            "reviewer_id": reviewer_id,
            "expert_slot": expert_id,
            "target_conflict": string(row["目标冲突代码"]),
            "evidence_sufficient": string(row["材料充分性代码"]),
            "need_more_context": string(row["更多上下文代码"]),
            "confidence": string(row["判断信心_1到5"]),
            "comments": string(row.get("整体备注")),
            "sections": by_case_sections.get(case_id, {}),
            "saved_at": string(row.get("保存时间")),
        }
    require(
        set(by_case_sections) == set(annotations),
        "XLSX section sheet has orphan cases or misses an entire case",
    )
    manifest_hashes = {string(row["manifest_hash"]) for row in case_rows} | {string(metadata["manifest_hash"])}
    translation_hashes = {string(row["translation_pack_hash"]) for row in case_rows} | {
        string(metadata["translation_pack_hash"])
    }
    require(len(manifest_hashes) == 1 and "" not in manifest_hashes, "XLSX manifest hashes differ/empty")
    require(len(translation_hashes) == 1 and "" not in translation_hashes, "XLSX translation hashes differ/empty")
    return {
        "schema_version": EXPORT_SCHEMA,
        "study_id": string(metadata["研究编号"]),
        "manifest_hash": next(iter(manifest_hashes)),
        "translation_pack_hash": next(iter(translation_hashes)),
        "annotation_mode": "blind_human_audit",
        "expert_slot": expert_id,
        "anonymous_reviewer_id": reviewer_id,
        "order": order,
        "active_index": len(order) - 1,
        "annotations": annotations,
        "exported_at": string(metadata.get("导出时间")),
    }


def load_expert_export(path: Path) -> dict[str, Any]:
    if path.suffix.lower() == ".json":
        payload = read_json(path)
        require(isinstance(payload, dict), f"Expert export {path} must be an object")
        return payload
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValidationError("Reading XLSX requires openpyxl>=3.1") from exc
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl exposes several file-format exceptions
            raise ValidationError(f"Cannot read XLSX {path}: {exc}") from exc
        try:
            required = {"样本整体标注", "逐段标注", "版本信息"}
            require(required <= set(workbook.sheetnames), f"XLSX {path} is missing sheets {sorted(required - set(workbook.sheetnames))}")
            return xlsx_tables_to_payload(
                _sheet_records(workbook["样本整体标注"]),
                _sheet_records(workbook["逐段标注"]),
                _sheet_records(workbook["版本信息"]),
            )
        finally:
            workbook.close()
    raise ValidationError(f"Unsupported export extension for {path}; use .json or .xlsx")


def validate_binary_annotation(
    annotation: Any, section_ids: set[str], *, label: str = "binary annotation"
) -> dict[str, Any]:
    """Validate one completed binary-form record without converting legacy labels.

    This is an independent data contract, not admission to the blind v2 pipeline.
    It preserves optional legacy fields verbatim and never manufactures role,
    need_more_context, or evidence_sufficient ratings.
    """
    require(isinstance(annotation, dict), f"{label}: annotation must be an object")
    require(annotation.get("annotation_protocol") == BINARY_PROTOCOL, f"{label}: binary protocol required")
    target = annotation.get("target_conflict")
    sufficient = annotation.get("judgment_sufficient")
    require(target in {"present", "absent"}, f"{label}: target_conflict must be present/absent")
    require(sufficient in {"yes", "no"}, f"{label}: judgment_sufficient must be yes/no")
    require(annotation.get("sets_reviewed") is True, f"{label}: material sets have not been confirmed")
    confidence = annotation.get("confidence")
    require(
        type(confidence) is int and 1 <= confidence <= 5
        or type(confidence) is str and confidence in {"1", "2", "3", "4", "5"},
        f"{label}: confidence must be an integer in 1..5",
    )
    sections = annotation.get("sections")
    require(isinstance(sections, dict) and set(sections) == section_ids, f"{label}: section coverage mismatch")
    for section_id, section in sections.items():
        require(isinstance(section, dict), f"{label}/{section_id}: section must be an object")
        require(
            type(section.get("minimal_evidence")) is bool
            and type(section.get("repair_packet")) is bool,
            f"{label}/{section_id}: set flags must be explicit booleans",
        )
    minimal = sorted(key for key, section in sections.items() if section["minimal_evidence"])
    repair = sorted(key for key, section in sections.items() if section["repair_packet"])
    require(set(minimal) <= set(repair), f"{label}: minimal evidence must be a subset of the repair packet")
    if target == "present" and sufficient == "yes":
        require(bool(minimal), f"{label}: sufficient positive requires minimal evidence")
    if target == "absent":
        require(not minimal, f"{label}: absent must not select minimal evidence")
    normalized = copy.deepcopy(annotation)
    normalized["confidence"] = int(confidence)
    normalized["minimal_evidence_section_ids"] = minimal
    normalized["repair_packet_section_ids"] = repair
    normalized["scoring_status"] = (
        "insufficient" if sufficient == "no" else "positive" if target == "present" else "negative"
    )
    return normalized


def binary_metric_applicability(annotation: Any, section_ids: set[str]) -> dict[str, Any]:
    """Return single-rating applicability, never a consensus or frozen gold label."""
    record = validate_binary_annotation(annotation, section_ids)
    status = record["scoring_status"]
    return {
        "annotation_protocol": BINARY_PROTOCOL,
        "target_conflict": record["target_conflict"],
        "judgment_sufficient": record["judgment_sufficient"],
        "scoring_status": status,
        "requires_adjudication": status == "insufficient",
        "positive_denominator_eligible": status == "positive",
        "negative_denominator_eligible": status == "negative",
        "false_positive_denominator_eligible": status == "negative",
        "cleanliness_denominator_eligible": status == "positive",
        "minimal_evidence_section_ids": record["minimal_evidence_section_ids"],
        "repair_packet_section_ids": record["repair_packet_section_ids"],
    }


def binary_case_metrics(
    annotation: Any,
    section_ids: set[str],
    returned_section_ids: Sequence[str],
    *,
    prediction: str,
    output_valid: bool,
    bridge_section_ids: Sequence[str] = (),
) -> dict[str, Any]:
    """Diagnostic single-record math; callers must bind case identity independently.

    Inputs are presentation-local section IDs, not raw node IDs. Bridges must be
    fixed independently of outputs. This does not create a formal label freeze.
    """
    eligibility = binary_metric_applicability(annotation, section_ids)
    returned = set(unique_strings(list(returned_section_ids), "binary.returned_section_ids"))
    bridges = set(unique_strings(list(bridge_section_ids), "binary.bridge_section_ids"))
    require(returned <= section_ids and bridges <= section_ids, "Binary metric IDs outside the case")
    require(type(output_valid) is bool, "Binary output_valid must be boolean")
    require(prediction in PREDICTIONS, "Binary prediction invalid")
    require(output_valid == (prediction != "invalid"), "Binary validity/prediction inconsistent")
    require(output_valid or not returned, "Binary invalid output must have empty returned set")
    evidence = set(eligibility["minimal_evidence_section_ids"])
    repair = set(eligibility["repair_packet_section_ids"])
    allowed = repair | bridges
    positive = eligibility["positive_denominator_eligible"]
    negative = eligibility["negative_denominator_eligible"]
    return {
        **eligibility,
        "analysis_scope": "single_record_diagnostic_not_consensus_or_frozen_gold",
        "output_valid": output_valid,
        "prediction": prediction,
        "returned_section_ids": sorted(returned),
        "bridge_section_ids": sorted(bridges),
        "strict_valid": float(output_valid),
        "positive_detection_hit": float(output_valid and prediction == "present") if positive else None,
        "negative_safe_hit": float(output_valid and prediction in {"absent", "insufficient", "abstain"}) if negative else None,
        "over_reported": float(output_valid and prediction == "present") if negative else None,
        "expert_evidence_recall": len(returned & evidence) / len(evidence) if positive else None,
        "repair_packet_coverage": len(returned & repair) / len(repair) if positive else None,
        "excess_node_rate": (len(returned - repair) / len(returned) if returned else 0.0) if positive else None,
        "bridge_aware_excess_node_rate": (len(returned - allowed) / len(returned) if returned else 0.0) if positive else None,
        "bridge_aware_precision": (len(returned & allowed) / len(returned) if returned else 0.0) if positive else None,
        "expert_evidence_density": (len(returned & evidence) / len(returned) if returned else 0.0) if positive else None,
        "empty_return_convention": "eligible positives retain denominator; recall/coverage/density/precision/excess are zero",
        "bridge_allowance_rule": "expert repair packet union independently fixed bridge section IDs",
    }


def binary_categorical_iaa(
    cases: Sequence[dict[str, Any]], field: str
) -> dict[str, Any]:
    """IAA for binary questions only; preserved legacy fields are not new ratings."""
    allowed = {
        "target_conflict": ("present", "absent"),
        "judgment_sufficient": ("yes", "no"),
        "role": (),
        "need_more_context": (),
        "evidence_sufficient": (),
    }
    require(field in allowed, f"Unsupported binary IAA field {field}")
    if not allowed[field]:
        return {"status": "not_collected_by_binary_protocol", "n_units": 0, "rating_count": 0}
    units = []
    excluded = 0
    for case in cases:
        annotations = case.get("annotations", {})
        require(isinstance(annotations, dict), "Binary IAA annotations must be an object")
        # Missing raters are not assigned a default value; incomplete units are reported.
        if any(expert not in annotations for expert in EXPERT_IDS):
            excluded += 1
            continue
        sections = set(case.get("section_ids", []))
        require(bool(sections), "Binary IAA case must declare section_ids")
        records = [validate_binary_annotation(annotations[expert], sections) for expert in EXPERT_IDS]
        units.append([record[field] for record in records])
    if not units:
        return {"status": "no_complete_three_rater_units", "n_units": 0, "rating_count": 0, "excluded_incomplete_units": excluded}
    return {
        **nominal_multi_rater(units, allowed[field]),
        "rating_count": len(units) * len(EXPERT_IDS),
        "excluded_incomplete_units": excluded,
        "field": field,
        "target_semantics": "raw reported judgment; not a sufficient-evidence gold label",
    }


def _validate_annotation(
    annotation: Any,
    expert_id: str,
    reviewer_id: str,
    case_id: str,
    section_ids: set[str],
) -> dict[str, Any]:
    require(isinstance(annotation, dict), f"{expert_id}/{case_id} annotation is not an object")
    require(
        not annotation.get("annotation_protocol"),
        f"{expert_id}/{case_id}: explicit annotation protocol cannot enter legacy blind v2 analysis",
    )
    record_reviewer = string(annotation.get("reviewer_id"))
    require(record_reviewer == reviewer_id, f"{expert_id}/{case_id} reviewer ID mismatch")
    if annotation.get("expert_slot") is not None:
        require(annotation.get("expert_slot") == expert_id, f"{expert_id}/{case_id} expert_slot mismatch")
    target = string(annotation.get("target_conflict"))
    sufficiency = string(annotation.get("evidence_sufficient"))
    context = string(annotation.get("need_more_context"))
    require(target in TARGET_LABELS, f"{expert_id}/{case_id} target_conflict invalid")
    require(sufficiency in SUFFICIENCY_LABELS, f"{expert_id}/{case_id} evidence_sufficient invalid")
    require(context in CONTEXT_LABELS, f"{expert_id}/{case_id} need_more_context invalid")
    try:
        confidence = int(annotation.get("confidence"))
    except (TypeError, ValueError) as exc:
        raise ValidationError(f"{expert_id}/{case_id} confidence invalid") from exc
    require(1 <= confidence <= 5, f"{expert_id}/{case_id} confidence outside 1..5")
    require(not (target == "absent" and sufficiency != "not_applicable"), f"{expert_id}/{case_id}: absent requires not_applicable")
    require(not (target == "present" and sufficiency == "not_applicable"), f"{expert_id}/{case_id}: present forbids not_applicable")
    sections = annotation.get("sections")
    require(isinstance(sections, dict) and set(sections) == section_ids, f"{expert_id}/{case_id} section coverage mismatch")
    normalized_sections: dict[str, dict[str, Any]] = {}
    for section_id, section in sections.items():
        require(isinstance(section, dict), f"{expert_id}/{case_id}/{section_id} is not an object")
        role = string(section.get("role"))
        require(role in SECTION_ROLES, f"{expert_id}/{case_id}/{section_id} role invalid")
        minimal = section.get("minimal_evidence")
        repair = section.get("repair_packet")
        require(type(minimal) is bool and type(repair) is bool, f"{expert_id}/{case_id}/{section_id} flags must be booleans")
        normalized_sections[section_id] = {
            "role": role,
            "minimal_evidence": minimal,
            "repair_packet": repair,
            "note": string(section.get("note")),
        }
    minimal_ids = sorted(key for key, value in normalized_sections.items() if value["minimal_evidence"])
    repair_ids = sorted(key for key, value in normalized_sections.items() if value["repair_packet"])
    require(
        set(minimal_ids) <= set(repair_ids),
        f"{expert_id}/{case_id}: minimal evidence must be a subset of the repair packet",
    )
    if target == "present" and sufficiency in {"yes", "partial"}:
        require(bool(minimal_ids), f"{expert_id}/{case_id} positive/sufficient case has no minimal evidence")
    if target == "absent":
        require(not minimal_ids, f"{expert_id}/{case_id} absent case must not select minimal evidence")
    return {
        "reviewer_id": reviewer_id,
        "target_conflict": target,
        "evidence_sufficient": sufficiency,
        "need_more_context": context,
        "confidence": confidence,
        "comments": string(annotation.get("comments")),
        "sections": normalized_sections,
        "minimal_evidence_section_ids": minimal_ids,
        "repair_packet_section_ids": repair_ids,
        "saved_at": string(annotation.get("saved_at")),
    }


def validate_expert_export(payload: dict[str, Any], layout: dict[str, Any], source: Path) -> dict[str, Any]:
    manifest = layout["manifest"]
    require(payload.get("schema_version") == EXPORT_SCHEMA, f"{source}: wrong export schema")
    require(is_synthetic(payload) == layout["synthetic_fixture"], f"{source}: synthetic marker mismatch")
    require(payload.get("study_id") == manifest["study_id"], f"{source}: study_id mismatch")
    require(payload.get("manifest_hash") == manifest["manifest_hash"], f"{source}: manifest_hash mismatch")
    require(payload.get("annotation_mode") == "blind_human_audit", f"{source}: annotation_mode mismatch")
    require(
        not payload.get("annotation_protocol"),
        f"{source}: explicit annotation protocol requires independent merge/adjudication/freeze; not legacy blind v2",
    )
    translation_hash = string(payload.get("translation_pack_hash"))
    require(translation_hash, f"{source}: empty translation_pack_hash")
    require(
        translation_hash == layout["translations"]["translation_pack_hash"],
        f"{source}: translation_pack_hash does not match the verified translation pack",
    )
    expert_id = string(payload.get("expert_slot"))
    require(expert_id in EXPERT_IDS, f"{source}: invalid expert_slot")
    reviewer_id = string(payload.get("anonymous_reviewer_id"))
    require(reviewer_id, f"{source}: anonymous_reviewer_id is empty")
    order = unique_strings(payload.get("order"), f"{source}: order")
    require(order == manifest["expert_orders"][expert_id], f"{source}: order differs from frozen expert order")
    annotations = payload.get("annotations")
    require(isinstance(annotations, dict), f"{source}: annotations must be an object")
    require(set(annotations) == set(layout["current"]), f"{source}: annotations are not exactly the current cases")
    public_by_id = {item["presentation_id"]: item for item in manifest["cases"]}
    normalized = {}
    for case_id in order:
        section_ids = {section["section_id"] for section in public_by_id[case_id]["sections"]}
        normalized[case_id] = _validate_annotation(
            annotations[case_id], expert_id, reviewer_id, case_id, section_ids
        )
    return {
        "expert_id": expert_id,
        "reviewer_id": reviewer_id,
        "translation_pack_hash": translation_hash,
        "order": order,
        "annotations": normalized,
    }


def majority(values: Sequence[str]) -> str | None:
    counts = Counter(values)
    if not counts:
        return None
    value, count = counts.most_common(1)[0]
    return value if count >= 2 else None


def majority_set(sets: Sequence[set[str]]) -> list[str]:
    counts = Counter(item for values in sets for item in values)
    return sorted(item for item, count in counts.items() if count >= 2)


def provisional_case(annotations: dict[str, dict[str, Any]]) -> dict[str, Any]:
    ordered = [annotations[expert_id] for expert_id in EXPERT_IDS]
    target = majority([item["target_conflict"] for item in ordered])
    sufficiency = majority([item["evidence_sufficient"] for item in ordered])
    context = majority([item["need_more_context"] for item in ordered])
    confidence_values = [item["confidence"] for item in ordered]
    minimal = majority_set([set(item["minimal_evidence_section_ids"]) for item in ordered])
    repair = majority_set([set(item["repair_packet_section_ids"]) for item in ordered])
    flags = []
    if target is None:
        flags.append("target_no_2_of_3_majority")
    if sufficiency is None:
        flags.append("sufficiency_no_2_of_3_majority")
    if context is None:
        flags.append("context_no_2_of_3_majority")
    if any(item["target_conflict"] == "insufficient" for item in ordered):
        flags.append("any_insufficient_context_label")
    if min(confidence_values) <= 2:
        flags.append("any_confidence_at_most_2")
    if target == "present" and sufficiency in {"yes", "partial"} and not minimal:
        flags.append("positive_without_2_of_3_minimal_evidence")
    if target == "present" and sufficiency in {"yes", "partial"} and not repair:
        flags.append("positive_without_2_of_3_repair_packet")
    return {
        "target_conflict": target,
        "evidence_sufficient": sufficiency,
        "need_more_context": context,
        "confidence_median": statistics.median(confidence_values),
        "minimal_evidence_section_ids": minimal,
        "repair_packet_section_ids": repair,
        "requires_adjudication": bool(flags),
        "adjudication_reasons": flags,
    }


def _pairwise_agreement(labels_a: Sequence[str], labels_b: Sequence[str]) -> dict[str, Any]:
    require(len(labels_a) == len(labels_b) and bool(labels_a), "Pairwise labels must be non-empty and aligned")
    observed = sum(a == b for a, b in zip(labels_a, labels_b)) / len(labels_a)
    categories = sorted(set(labels_a) | set(labels_b))
    count_a = Counter(labels_a)
    count_b = Counter(labels_b)
    n = len(labels_a)
    expected = sum((count_a[category] / n) * (count_b[category] / n) for category in categories)
    kappa = None if math.isclose(1.0 - expected, 0.0) else (observed - expected) / (1.0 - expected)
    return {"n_units": n, "raw_agreement": observed, "cohen_kappa": kappa}


def nominal_multi_rater(
    units: Sequence[Sequence[str]], category_universe: Sequence[str]
) -> dict[str, Any]:
    require(bool(units), "Nominal agreement needs at least one unit")
    categories = unique_strings(list(category_universe), "category_universe")
    require(len(categories) >= 2, "Nominal agreement needs a fixed universe of at least two categories")
    category_set = set(categories)
    pair_agreements: list[float] = []
    pooled = Counter()
    disagreements = 0
    pair_count = 0
    unanimous = 0
    for labels in units:
        require(len(labels) == len(EXPERT_IDS), "Each IAA unit must have exactly three labels")
        require(set(labels) <= category_set, "IAA label is outside the fixed category universe")
        pooled.update(labels)
        unanimous += len(set(labels)) == 1
        for left, right in itertools.combinations(labels, 2):
            pair_agreements.append(float(left == right))
            disagreements += left != right
            pair_count += 1
    observed_agreement = sum(pair_agreements) / len(pair_agreements)
    observed_disagreement = disagreements / pair_count
    total = sum(pooled.values())
    expected_disagreement = 1.0 - sum(count * (count - 1) for count in pooled.values()) / (total * (total - 1))
    alpha = None if math.isclose(expected_disagreement, 0.0) else 1.0 - observed_disagreement / expected_disagreement
    category_count = len(categories)
    chance_ac1 = sum(
        (pooled[category] / total) * (1.0 - pooled[category] / total)
        for category in categories
    ) / (category_count - 1)
    ac1 = None if math.isclose(1.0 - chance_ac1, 0.0) else (observed_agreement - chance_ac1) / (1.0 - chance_ac1)
    return {
        "n_units": len(units),
        "ratings_per_unit": len(EXPERT_IDS),
        "unanimous_unit_rate": unanimous / len(units),
        "mean_pairwise_raw_agreement": observed_agreement,
        "krippendorff_alpha_nominal": alpha,
        "gwet_ac1_nominal": ac1,
        "fixed_category_universe": categories,
        "label_counts": {category: pooled[category] for category in categories},
    }


def categorical_iaa(cases: Sequence[dict[str, Any]], field: str, section_level: bool = False) -> dict[str, Any]:
    units: list[list[str]] = []
    per_expert: dict[str, list[str]] = {expert_id: [] for expert_id in EXPERT_IDS}
    if section_level:
        for case in cases:
            for section_id in sorted(case["section_id_to_node_id"]):
                labels = [case["annotations"][expert_id]["sections"][section_id][field] for expert_id in EXPERT_IDS]
                units.append(labels)
                for expert_id, label in zip(EXPERT_IDS, labels):
                    per_expert[expert_id].append(label)
    else:
        for case in cases:
            labels = [case["annotations"][expert_id][field] for expert_id in EXPERT_IDS]
            units.append(labels)
            for expert_id, label in zip(EXPERT_IDS, labels):
                per_expert[expert_id].append(label)
    category_universes = {
        "target_conflict": TARGET_LABELS,
        "evidence_sufficient": SUFFICIENCY_LABELS,
        "need_more_context": CONTEXT_LABELS,
        "role": SECTION_ROLES,
    }
    require(field in category_universes, f"No frozen category universe for IAA field {field}")
    result = nominal_multi_rater(units, category_universes[field])
    result["pairwise"] = {
        f"{left}__{right}": _pairwise_agreement(per_expert[left], per_expert[right])
        for left, right in itertools.combinations(EXPERT_IDS, 2)
    }
    return result


def jaccard(left: set[str], right: set[str]) -> float:
    if not left and not right:
        return 1.0
    return len(left & right) / len(left | right)


def set_iaa(cases: Sequence[dict[str, Any]], field: str) -> dict[str, Any]:
    pair_values: dict[str, list[float]] = {
        f"{left}__{right}": [] for left, right in itertools.combinations(EXPERT_IDS, 2)
    }
    pair_eligible_values: dict[str, list[float]] = {
        pair: [] for pair in pair_values
    }
    for case in cases:
        for left, right in itertools.combinations(EXPERT_IDS, 2):
            left_set = set(case["annotations"][left][field])
            right_set = set(case["annotations"][right][field])
            pair = f"{left}__{right}"
            value = jaccard(left_set, right_set)
            pair_values[pair].append(value)
            if left_set or right_set:
                pair_eligible_values[pair].append(value)
    flattened = [value for values in pair_values.values() for value in values]
    eligible_flattened = [value for values in pair_eligible_values.values() for value in values]
    return {
        "n_cases": len(cases),
        "empty_empty_jaccard_convention": 1.0,
        "mean_pairwise_jaccard": sum(flattened) / len(flattened),
        "pairwise_mean_jaccard": {
            pair: sum(values) / len(values) for pair, values in pair_values.items()
        },
        "eligible_definition": "pair/case union is non-empty; empty-empty pairs are excluded",
        "eligible_pair_case_count": len(eligible_flattened),
        "eligible_conditional_mean_pairwise_jaccard": (
            sum(eligible_flattened) / len(eligible_flattened)
            if eligible_flattened else None
        ),
        "pairwise_eligible_conditional": {
            pair: {
                "eligible_case_count": len(values),
                "mean_jaccard": (sum(values) / len(values) if values else None),
            }
            for pair, values in pair_eligible_values.items()
        },
    }


def build_iaa(merged: dict[str, Any]) -> dict[str, Any]:
    strata: dict[str, Any] = {}
    groups = {
        "pre18": [case for case in merged["cases"] if case["stratum"] == "pre18"],
        "post_new24": [case for case in merged["cases"] if case["stratum"] == "post_new24"],
    }
    for stratum, cases in groups.items():
        require(bool(cases), f"IAA stratum {stratum} is empty")
        strata[stratum] = {
            "n_cases": len(cases),
            "target_conflict": categorical_iaa(cases, "target_conflict"),
            "evidence_sufficient": categorical_iaa(cases, "evidence_sufficient"),
            "need_more_context": categorical_iaa(cases, "need_more_context"),
            "section_role": categorical_iaa(cases, "role", section_level=True),
            "minimal_evidence": set_iaa(cases, "minimal_evidence_section_ids"),
            "repair_packet": set_iaa(cases, "repair_packet_section_ids"),
        }
    return seal({
        "schema_version": IAA_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": merged["synthetic_fixture"],
        "merged_artifact_sha256": merged["artifact_sha256"],
        "unit_warning": "Cases/source SCCs are units; three experts are raters, not independent samples.",
        "legacy_policy": "post_legacy12 is not recomputed or pooled with the current blind round.",
        "strata": strata,
    })


def build_adjudication_template(merged: dict[str, Any]) -> dict[str, Any]:
    cases = []
    for case in merged["cases"]:
        provisional = case["provisional"]
        if not provisional["requires_adjudication"]:
            continue
        cases.append({
            "case_id": case["case_id"],
            "stratum": case["stratum"],
            "domain": case["domain"],
            "required_reasons": provisional["adjudication_reasons"],
            "allowed_section_ids": sorted(case["section_id_to_node_id"]),
            "provisional_2_of_3": provisional,
            "decision": {
                "final_target_conflict": None,
                "final_evidence_sufficient": None,
                "final_need_more_context": None,
                "final_minimal_evidence_section_ids": None,
                "final_repair_packet_section_ids": None,
                "adjudicator_id": None,
                "adjudication_note": None,
            },
        })
    return seal({
        "schema_version": ADJUDICATION_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": merged["synthetic_fixture"],
        "merged_artifact_sha256": merged["artifact_sha256"],
        "selection_rule": ADJUDICATION_SELECTION_RULE,
        "system_outcomes_must_remain_unseen": True,
        "outcome_blind_attestation": {
            "system_outcomes_unseen": True,
            "attested_by": None,
            "attested_at": None,
        },
        "cases": cases,
    })


def merge_exports(
    manifest_path: Path,
    codebook_path: Path,
    translations_path: Path,
    export_paths: Sequence[Path],
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    require(len(export_paths) == 3, "Exactly three expert exports are required")
    layout = load_layout(manifest_path, codebook_path, translations_path)
    experts: dict[str, dict[str, Any]] = {}
    sources = []
    for path in export_paths:
        validated = validate_expert_export(load_expert_export(path), layout, path)
        expert_id = validated["expert_id"]
        require(expert_id not in experts, f"Duplicate expert slot {expert_id}")
        experts[expert_id] = validated
        sources.append({"path": str(path.resolve()), "file_sha256": file_sha256(path), "expert_id": expert_id})
    require(set(experts) == set(EXPERT_IDS), "Exports must contain expert_1, expert_2, expert_3 exactly once")
    reviewer_ids = [experts[expert_id]["reviewer_id"] for expert_id in EXPERT_IDS]
    require(len(set(reviewer_ids)) == 3, "Anonymous reviewer IDs must be distinct")
    translation_hashes = {experts[expert_id]["translation_pack_hash"] for expert_id in EXPERT_IDS}
    require(len(translation_hashes) == 1, "Expert exports use different translation packs")

    cases = []
    manifest_order = [item["presentation_id"] for item in layout["manifest"]["cases"]]
    for case_id in manifest_order:
        identity = layout["current"][case_id]
        annotations = {
            expert_id: experts[expert_id]["annotations"][case_id] for expert_id in EXPERT_IDS
        }
        cases.append({**identity, "annotations": annotations, "provisional": provisional_case(annotations)})
    counts = Counter(case["stratum"] for case in cases)
    merged = seal({
        "schema_version": MERGED_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": layout["synthetic_fixture"],
        "study_id": layout["manifest"]["study_id"],
        "manifest_hash": layout["manifest"]["manifest_hash"],
        "translation_pack_hash": next(iter(translation_hashes)),
        "source_identity": layout["source_identity"],
        "source_exports": sorted(sources, key=lambda row: row["expert_id"]),
        "expert_slots": list(EXPERT_IDS),
        "reviewer_ids": {expert_id: experts[expert_id]["reviewer_id"] for expert_id in EXPERT_IDS},
        "counts": dict(counts),
        "legacy_case_count_registered_but_not_in_iaa": len(layout["legacy"]),
        "cases": cases,
    })
    return merged, build_iaa(merged), build_adjudication_template(merged)


def merge_to_dir(
    manifest_path: Path,
    codebook_path: Path,
    translations_path: Path,
    export_paths: Sequence[Path],
    out_dir: Path,
) -> dict[str, Path]:
    merged, iaa, adjudication = merge_exports(
        manifest_path, codebook_path, translations_path, export_paths
    )
    paths = {
        "merged": out_dir / "merged_bundle.json",
        "iaa": out_dir / "iaa.json",
        "adjudication": out_dir / "adjudication_template.json",
    }
    write_json(paths["merged"], merged)
    write_json(paths["iaa"], iaa)
    write_json(paths["adjudication"], adjudication)
    return paths


def _validate_final_fields(
    target: Any,
    sufficiency: Any,
    context: Any,
    minimal: Any,
    repair: Any,
    allowed_ids: set[str],
    label: str,
) -> tuple[str, str, str, list[str], list[str]]:
    target_s = string(target)
    sufficiency_s = string(sufficiency)
    context_s = string(context)
    require(target_s in TARGET_LABELS, f"{label}: final_target_conflict invalid")
    require(sufficiency_s in SUFFICIENCY_LABELS, f"{label}: final_evidence_sufficient invalid")
    require(context_s in CONTEXT_LABELS, f"{label}: final_need_more_context invalid")
    minimal_ids = unique_strings(minimal, f"{label}.minimal_evidence")
    repair_ids = unique_strings(repair, f"{label}.repair_packet")
    require(set(minimal_ids) <= allowed_ids, f"{label}: unknown minimal-evidence IDs")
    require(set(repair_ids) <= allowed_ids, f"{label}: unknown repair-packet IDs")
    require(not (target_s == "absent" and sufficiency_s != "not_applicable"), f"{label}: absent requires not_applicable")
    require(not (target_s == "present" and sufficiency_s == "not_applicable"), f"{label}: present forbids not_applicable")
    if target_s == "present" and sufficiency_s in {"yes", "partial"}:
        require(bool(minimal_ids), f"{label}: positive/sufficient final label needs minimal evidence")
        require(bool(repair_ids), f"{label}: positive/sufficient final label needs a repair packet")
    if target_s == "absent":
        require(not minimal_ids, f"{label}: absent final label must have no minimal evidence")
    require(
        set(minimal_ids) <= set(repair_ids),
        f"{label}: minimal evidence must be a subset of the repair packet",
    )
    return target_s, sufficiency_s, context_s, sorted(minimal_ids), sorted(repair_ids)


def load_legacy_labels(path: Path, layout: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    require(
        path.is_file(),
        f"Formal legacy12 label gate BLOCKED: explicit verified legacy label artifact is missing: {path}",
    )
    payload = read_json(path)
    require(isinstance(payload, dict) and payload.get("schema_version") == LEGACY_SCHEMA, "Wrong legacy label schema")
    require(is_synthetic(payload) == layout["synthetic_fixture"], "Legacy synthetic marker mismatch")
    declared_source_hash = require_sha256(
        payload.get("source_artifact_sha256"), "Legacy source_artifact_sha256"
    )
    source_artifact_path = _resolve_linked_path(
        payload.get("source_artifact_path"), path, "Legacy source artifact"
    )
    require(source_artifact_path != path.resolve(), "Legacy source artifact cannot be the normalized label file itself")
    actual_source_hash = file_sha256(source_artifact_path)
    require(
        declared_source_hash == actual_source_hash,
        f"Legacy source artifact hash mismatch: declared {declared_source_hash}, actual {actual_source_hash}",
    )
    source_artifact = read_json(source_artifact_path)
    require(
        isinstance(source_artifact, dict)
        and source_artifact.get("schema_version") == LEGACY_SOURCE_SCHEMA,
        "Legacy source artifact must use the supported formal source-record schema",
    )
    require(
        is_synthetic(source_artifact) == layout["synthetic_fixture"],
        "Legacy source artifact synthetic marker mismatch",
    )
    raw_source_records = source_artifact.get("records")
    require(isinstance(raw_source_records, list), "Legacy source artifact records must be a list")
    source_by_locator: dict[str, dict[str, Any]] = {}
    for source_record in raw_source_records:
        require(isinstance(source_record, dict), "Legacy source record is not an object")
        locator = string(source_record.get("source_record_locator"))
        require(locator and locator not in source_by_locator, "Legacy source record locators empty/duplicate")
        source_by_locator[locator] = source_record
    require(bool(string(payload.get("transcription_certified_by"))), "Legacy transcription must be certified")
    validate_iso_datetime(string(payload.get("transcription_certified_at")), "legacy transcription_certified_at")
    raw_cases = payload.get("cases")
    require(isinstance(raw_cases, list), "Legacy cases must be a list")
    by_id = {string(item.get("case_id")): item for item in raw_cases if isinstance(item, dict)}
    require(len(by_id) == len(raw_cases), "Legacy case IDs are empty/duplicate")
    require(set(by_id) == set(layout["legacy"]), "Legacy labels do not exactly cover registered old cases")
    result = []
    source_locators: set[str] = set()
    for case_id, identity in layout["legacy"].items():
        raw = by_id[case_id]
        require(string(raw.get("legacy_case_id")) == identity["legacy_case_id"], f"Legacy {case_id} legacy_case_id mismatch")
        require(
            string(raw.get("post_content_fingerprint")) == identity["post_content_fingerprint"],
            f"Legacy {case_id} post_content_fingerprint mismatch",
        )
        require(string(raw.get("source_scc_uid")) == identity["source_scc_uid"], f"Legacy {case_id} SCC mismatch")
        require(string(raw.get("domain")) == identity["domain"], f"Legacy {case_id} domain mismatch")
        source_locator = string(raw.get("source_record_locator"))
        provenance = string(raw.get("label_provenance"))
        require(source_locator and source_locator not in source_locators, f"Legacy {case_id} source_record_locator empty/duplicate")
        require(provenance, f"Legacy {case_id} label_provenance is required")
        require(source_locator in source_by_locator, f"Legacy {case_id} source_record_locator does not resolve")
        source_locators.add(source_locator)
        target, sufficiency, context, minimal, repair = _validate_final_fields(
            raw.get("final_target_conflict"),
            raw.get("final_evidence_sufficient"),
            raw.get("final_need_more_context"),
            raw.get("minimal_evidence_node_ids"),
            raw.get("repair_packet_node_ids"),
            set(identity["node_universe"]),
            f"legacy/{case_id}",
        )
        endpoints = unique_strings(raw.get("endpoint_node_ids", []), f"legacy/{case_id}.endpoint_node_ids")
        require(set(endpoints) <= set(identity["node_universe"]), f"Legacy {case_id} has unknown endpoints")
        if target == "present":
            require(bool(endpoints), f"Legacy positive {case_id} must bind its injected endpoint nodes")
        require(
            set(endpoints) == set(identity["endpoint_node_ids"]),
            f"Legacy {case_id} endpoints differ from the canonical injection identity",
        )
        source_record = source_by_locator[source_locator]
        source_target, source_sufficiency, source_context, source_minimal, source_repair = _validate_final_fields(
            source_record.get("final_target_conflict"),
            source_record.get("final_evidence_sufficient"),
            source_record.get("final_need_more_context"),
            source_record.get("minimal_evidence_node_ids"),
            source_record.get("repair_packet_node_ids"),
            set(identity["node_universe"]),
            f"legacy-source/{source_locator}",
        )
        source_endpoints = unique_strings(
            source_record.get("endpoint_node_ids", []),
            f"legacy-source/{source_locator}.endpoint_node_ids",
        )
        normalized_projection = {
            "case_id": case_id,
            "legacy_case_id": identity["legacy_case_id"],
            "post_content_fingerprint": identity["post_content_fingerprint"],
            "source_scc_uid": identity["source_scc_uid"],
            "domain": identity["domain"],
            "final_target_conflict": target,
            "final_evidence_sufficient": sufficiency,
            "final_need_more_context": context,
            "minimal_evidence_node_ids": minimal,
            "repair_packet_node_ids": repair,
            "endpoint_node_ids": sorted(endpoints),
            "label_provenance": provenance,
        }
        source_projection = {
            "case_id": string(source_record.get("case_id")),
            "legacy_case_id": string(source_record.get("legacy_case_id")),
            "post_content_fingerprint": string(source_record.get("post_content_fingerprint")),
            "source_scc_uid": string(source_record.get("source_scc_uid")),
            "domain": string(source_record.get("domain")),
            "final_target_conflict": source_target,
            "final_evidence_sufficient": source_sufficiency,
            "final_need_more_context": source_context,
            "minimal_evidence_node_ids": source_minimal,
            "repair_packet_node_ids": source_repair,
            "endpoint_node_ids": sorted(source_endpoints),
            "label_provenance": string(source_record.get("label_provenance")),
        }
        require(
            source_projection == normalized_projection,
            f"Legacy {case_id} normalized row differs from resolved source record {source_locator}",
        )
        result.append({
            **identity,
            "final_target_conflict": target,
            "final_evidence_sufficient": sufficiency,
            "final_need_more_context": context,
            "minimal_evidence_node_ids": minimal,
            "repair_packet_node_ids": repair,
            "endpoint_node_ids": sorted(endpoints),
            "bridge_node_ids": identity["bridge_node_ids"],
            "decision_basis": "verbatim_legacy_label_reuse",
            "legacy_label_provenance": provenance,
            "legacy_source_record_locator": source_locator,
        })
    require(
        source_locators == set(source_by_locator),
        "Legacy source artifact records are not a one-to-one match for normalized locators",
    )
    payload = copy.deepcopy(payload)
    payload["verified_source_artifact_path"] = str(source_artifact_path)
    payload["verified_source_artifact_sha256"] = actual_source_hash
    return result, payload


def freeze_labels(
    merged_path: Path,
    adjudication_path: Path,
    manifest_path: Path,
    codebook_path: Path,
    translations_path: Path,
    legacy_path: Path,
    *,
    export_paths: Sequence[Path],
    protocol_id: str,
    frozen_by: str,
    frozen_at: str,
    confirm_outcome_blind_freeze: bool,
    score_models: Sequence[str] | None = None,
    score_methods: Sequence[str] | None = None,
    score_strata: Sequence[str] | None = None,
) -> dict[str, Any]:
    require(confirm_outcome_blind_freeze, "Freeze requires --confirm-outcome-blind-freeze attestation")
    require(bool(string(protocol_id)), "protocol_id is required")
    require(bool(string(frozen_by)), "frozen_by is required")
    validate_iso_datetime(frozen_at, "frozen_at")
    frozen_models = unique_strings(list(score_models or []), "score_models")
    frozen_methods = unique_strings(list(score_methods or []), "score_methods")
    frozen_strata = unique_strings(list(score_strata or []), "score_strata")
    require(frozen_models and frozen_methods, "Freeze requires non-empty score model and method axes")
    require(
        set(frozen_strata) == set(FORMAL_STRATUM_COUNTS),
        "Freeze score_strata must contain all three predeclared strata exactly",
    )
    merged = read_json(merged_path)
    require(isinstance(merged, dict) and merged.get("schema_version") == MERGED_SCHEMA, "Wrong merged schema")
    verify_seal(merged)
    layout = load_layout(manifest_path, codebook_path, translations_path)
    require(merged.get("manifest_hash") == layout["manifest"]["manifest_hash"], "Merged/layout manifest mismatch")
    require(is_synthetic(merged) == layout["synthetic_fixture"], "Merged synthetic marker mismatch")
    require(
        merged.get("source_identity") == layout["source_identity"],
        "Merged source identity no longer matches manifest/codebook/source_master/translation inputs",
    )
    bound_exports = merged.get("source_exports")
    require(isinstance(bound_exports, list) and len(bound_exports) == 3, "Merged bundle must bind exactly three source exports")
    bound_by_expert: dict[str, Path] = {}
    for row in bound_exports:
        require(isinstance(row, dict), "Merged source_exports row is not an object")
        expert_id = string(row.get("expert_id"))
        require(expert_id in EXPERT_IDS and expert_id not in bound_by_expert, "Merged source expert IDs invalid/duplicate")
        source_path = Path(string(row.get("path"))).resolve()
        require(source_path.is_file(), f"Bound source export is missing: {source_path}")
        require(file_sha256(source_path) == row.get("file_sha256"), f"Bound source export hash mismatch for {expert_id}")
        bound_by_expert[expert_id] = source_path
    require(set(bound_by_expert) == set(EXPERT_IDS), "Merged source exports do not cover all three expert slots")
    require(len(export_paths) == 3, "Freeze requires exactly three explicit source exports")
    supplied = {path.resolve() for path in export_paths}
    require(
        supplied == set(bound_by_expert.values()),
        "Freeze --exports do not exactly match the three source-bound merged exports",
    )
    recomputed_merged, recomputed_iaa, recomputed_queue = merge_exports(
        manifest_path,
        codebook_path,
        translations_path,
        [bound_by_expert[expert_id] for expert_id in EXPERT_IDS],
    )
    require(
        recomputed_merged["artifact_sha256"] == merged["artifact_sha256"],
        "Merged bundle differs from a fresh recomputation of its three source exports",
    )
    # From here onward, every provisional label and queue decision is sourced from
    # the fresh recomputation rather than the caller-supplied merged body.
    merged = recomputed_merged
    merged_cases = {case["case_id"]: case for case in merged.get("cases", []) if isinstance(case, dict)}
    require(len(merged_cases) == len(merged.get("cases", [])), "Merged case IDs empty/duplicate")
    require(set(merged_cases) == set(layout["current"]), "Merged current cases differ from current layout")
    identity_fields = (
        "original_case_id",
        "case_uid",
        "stratum",
        "domain",
        "source_scc_uid",
        "post_content_fingerprint",
        "section_id_to_node_id",
        "node_universe",
        "endpoint_node_ids",
        "bridge_node_ids",
        "post_identity_record_sha256",
    )
    for case_id, expected in layout["current"].items():
        require(
            all(merged_cases[case_id].get(field) == expected[field] for field in identity_fields),
            f"Merged identity mapping differs from verified codebook for {case_id}",
        )
    adjudication = read_json(adjudication_path)
    require(isinstance(adjudication, dict) and adjudication.get("schema_version") == ADJUDICATION_SCHEMA, "Wrong adjudication schema")
    verify_seal(adjudication)
    require(adjudication.get("merged_artifact_sha256") == merged["artifact_sha256"], "Adjudication/merged hash mismatch")
    require(is_synthetic(adjudication) == layout["synthetic_fixture"], "Adjudication synthetic marker mismatch")
    require(
        adjudication.get("system_outcomes_must_remain_unseen") is True,
        "Adjudication must state system_outcomes_must_remain_unseen=true",
    )
    require(
        adjudication.get("selection_rule") == ADJUDICATION_SELECTION_RULE,
        "Adjudication selection_rule differs from the frozen rule",
    )
    outcome_attestation = adjudication.get("outcome_blind_attestation")
    require(isinstance(outcome_attestation, dict), "Adjudication lacks outcome_blind_attestation")
    require(
        outcome_attestation.get("system_outcomes_unseen") is True,
        "Outcome-blind attestation must explicitly state system_outcomes_unseen=true",
    )
    require(bool(string(outcome_attestation.get("attested_by"))), "Outcome-blind attested_by is required")
    validate_iso_datetime(string(outcome_attestation.get("attested_at")), "outcome-blind attested_at")
    raw_decisions = adjudication.get("cases")
    require(isinstance(raw_decisions, list), "Adjudication cases must be a list")
    decisions = {string(row.get("case_id")): row for row in raw_decisions if isinstance(row, dict)}
    require(len(decisions) == len(raw_decisions), "Adjudication case IDs empty/duplicate")
    frozen_current = []
    recomputed_queue_by_id = {
        row["case_id"]: row for row in recomputed_queue["cases"]
    }
    required_ids = set(recomputed_queue_by_id)
    require(
        set(decisions) == required_ids,
        f"Adjudication case set must exactly match the predeclared queue; missing={sorted(required_ids - set(decisions))}, extra={sorted(set(decisions) - required_ids)}",
    )
    for case in merged["cases"]:
        case_id = case["case_id"]
        provisional = case["provisional"]
        raw_row = decisions.get(case_id)
        if raw_row is not None:
            expected_queue_row = recomputed_queue_by_id[case_id]
            for metadata_field in (
                "stratum",
                "domain",
                "required_reasons",
                "allowed_section_ids",
                "provisional_2_of_3",
            ):
                require(
                    raw_row.get(metadata_field) == expected_queue_row[metadata_field],
                    f"{case_id}: adjudication metadata {metadata_field} differs from recomputed queue",
                )
            require(
                raw_row.get("required_reasons") == provisional["adjudication_reasons"],
                f"{case_id}: adjudication reasons differ from predeclared queue",
            )
            require(
                raw_row.get("allowed_section_ids") == sorted(case["section_id_to_node_id"]),
                f"{case_id}: adjudication allowed_section_ids differ from merged case",
            )
        raw_decision = raw_row.get("decision") if raw_row else None
        adjudicated = isinstance(raw_decision, dict) and raw_decision.get("final_target_conflict") is not None
        if provisional["requires_adjudication"]:
            require(adjudicated, f"Required adjudication {case_id} is not filled")
        if adjudicated:
            require(bool(string(raw_decision.get("adjudicator_id"))), f"{case_id}: adjudicator_id required")
            require(bool(string(raw_decision.get("adjudication_note"))), f"{case_id}: adjudication_note required")
            values = (
                raw_decision.get("final_target_conflict"),
                raw_decision.get("final_evidence_sufficient"),
                raw_decision.get("final_need_more_context"),
                raw_decision.get("final_minimal_evidence_section_ids"),
                raw_decision.get("final_repair_packet_section_ids"),
            )
            basis = "domain_adjudication"
        else:
            values = (
                provisional.get("target_conflict"),
                provisional.get("evidence_sufficient"),
                provisional.get("need_more_context"),
                provisional.get("minimal_evidence_section_ids"),
                provisional.get("repair_packet_section_ids"),
            )
            basis = "predeclared_2_of_3_majority"
        target, sufficiency, context, minimal_sections, repair_sections = _validate_final_fields(
            *values,
            set(case["section_id_to_node_id"]),
            f"current/{case_id}",
        )
        mapping = case["section_id_to_node_id"]
        frozen_current.append({
            "case_id": case_id,
            "original_case_id": case["original_case_id"],
            "case_uid": case["case_uid"],
            "stratum": case["stratum"],
            "domain": case["domain"],
            "source_scc_uid": case["source_scc_uid"],
            "post_content_fingerprint": case["post_content_fingerprint"],
            "post_identity_record_sha256": case["post_identity_record_sha256"],
            "node_universe": case["node_universe"],
            "final_target_conflict": target,
            "final_evidence_sufficient": sufficiency,
            "final_need_more_context": context,
            "confidence_median": provisional["confidence_median"],
            "minimal_evidence_section_ids": minimal_sections,
            "repair_packet_section_ids": repair_sections,
            "minimal_evidence_node_ids": sorted(mapping[item] for item in minimal_sections),
            "repair_packet_node_ids": sorted(mapping[item] for item in repair_sections),
            "endpoint_node_ids": case["endpoint_node_ids"],
            "bridge_node_ids": case["bridge_node_ids"],
            "bridge_aware_allowed_node_ids": sorted(
                set(mapping[item] for item in repair_sections) | set(case["bridge_node_ids"])
            ),
            "endpoint_metric_available": bool(case["endpoint_node_ids"]),
            "decision_basis": basis,
        })

    legacy_cases, legacy_payload = load_legacy_labels(legacy_path, layout)
    for case in legacy_cases:
        case["bridge_aware_allowed_node_ids"] = sorted(
            set(case["repair_packet_node_ids"]) | set(case["bridge_node_ids"])
        )
        case["endpoint_metric_available"] = bool(case["endpoint_node_ids"])
    all_cases = frozen_current + legacy_cases
    counts = Counter(case["stratum"] for case in all_cases)
    if not layout["synthetic_fixture"]:
        require(dict(counts) == FORMAL_STRATUM_COUNTS, "Frozen formal stratum counts mismatch")
    endpoint_metric_scope = []
    for stratum in FORMAL_STRATUM_COUNTS:
        availability = {
            bool(case["endpoint_metric_available"])
            for case in all_cases
            if case["stratum"] == stratum
        }
        require(len(availability) == 1, f"Endpoint availability is inconsistent within stratum {stratum}")
        if availability == {True}:
            endpoint_metric_scope.append(stratum)
    payload = {
        "schema_version": FROZEN_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": layout["synthetic_fixture"],
        "study_id": merged["study_id"],
        "manifest_hash": merged["manifest_hash"],
        "source_identity": layout["source_identity"],
        "protocol_id": protocol_id,
        "freeze_attestation": {
            "frozen_by": frozen_by,
            "frozen_at": frozen_at,
            "system_outcomes_unseen": outcome_attestation["system_outcomes_unseen"],
            "outcome_blind_attested_by": outcome_attestation["attested_by"],
            "outcome_blind_attested_at": outcome_attestation["attested_at"],
        },
        "source_artifacts": {
            "merged_artifact_sha256": merged["artifact_sha256"],
            "merged_file_sha256": file_sha256(merged_path),
            "adjudication_file_sha256": file_sha256(adjudication_path),
            "legacy_file_sha256": file_sha256(legacy_path),
            "legacy_source_artifact_sha256": legacy_payload["verified_source_artifact_sha256"],
            "legacy_source_artifact_path": legacy_payload["verified_source_artifact_path"],
            "source_exports": merged["source_exports"],
            "recomputed_iaa_artifact_sha256": recomputed_iaa["artifact_sha256"],
            "recomputed_adjudication_queue_artifact_sha256": recomputed_queue["artifact_sha256"],
        },
        "score_protocol": {
            "models": frozen_models,
            "methods": frozen_methods,
            "strata": frozen_strata,
            "require_complete_grid": True,
            "endpoint_metric_scope": endpoint_metric_scope,
            "endpoint_availability_rule": "report only strata where every frozen case has canonical non-empty endpoint IDs",
            "bridge_aware_cleanliness": {
                "enabled": True,
                "allowed_node_rule": "frozen repair packet union canonical injection bridge nodes",
                "excess_denominator": "returned node count; zero for empty returned set",
            },
        },
        "stratum_policy": {
            "pre18": "current blind round",
            "post_new24": "current blind round",
            "post_legacy12": "verbatim legacy reuse; separate reporting only",
            "pool_old_and_new_post": False,
            "endpoint_metric_scope": endpoint_metric_scope,
        },
        "counts": dict(counts),
        "cases": all_cases,
    }
    return seal(payload, "label_bundle_sha256")


def percentile(values: Sequence[float], probability: float) -> float:
    require(bool(values), "Percentile needs values")
    ordered = sorted(values)
    index = (len(ordered) - 1) * probability
    left = math.floor(index)
    right = math.ceil(index)
    if left == right:
        return ordered[left]
    return ordered[left] * (right - index) + ordered[right] * (index - left)


def cluster_bootstrap_ci(
    rows: Sequence[dict[str, Any]],
    metric: str,
    *,
    samples: int,
    seed: int,
) -> tuple[float, float, float, int, int]:
    require(samples > 0, "Bootstrap samples must be positive")
    eligible = [row for row in rows if row.get(metric) is not None]
    require(bool(eligible), f"No eligible rows for metric {metric}")
    clusters: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in eligible:
        cluster = string(row.get("source_scc_uid"))
        require(cluster, "Clustered row has no source_scc_uid")
        clusters[cluster].append(row)
    keys = sorted(clusters)
    cluster_means = {
        key: sum(float(row[metric]) for row in cluster_rows) / len(cluster_rows)
        for key, cluster_rows in clusters.items()
    }
    point = sum(cluster_means.values()) / len(cluster_means)
    rng = random.Random(seed)
    estimates = []
    for _ in range(samples):
        sampled_keys = rng.choices(keys, k=len(keys))
        estimates.append(sum(cluster_means[key] for key in sampled_keys) / len(sampled_keys))
    return point, percentile(estimates, 0.025), percentile(estimates, 0.975), len(eligible), len(keys)


def _stable_group_seed(base_seed: int, parts: Iterable[str]) -> int:
    digest = hashlib.sha256((str(base_seed) + "|" + "|".join(parts)).encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big")


def score_system_outputs(
    frozen_path: Path,
    system_outputs_path: Path,
    *,
    bootstrap_samples: int,
    seed: int,
    require_complete_grid: bool,
    expected_models: Sequence[str] | None = None,
    expected_methods: Sequence[str] | None = None,
    expected_strata: Sequence[str] | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    frozen = read_json(frozen_path)
    require(isinstance(frozen, dict) and frozen.get("schema_version") == FROZEN_SCHEMA, "Wrong frozen-label schema")
    verify_seal(frozen, "label_bundle_sha256")
    system = read_json(system_outputs_path)
    require(isinstance(system, dict) and system.get("schema_version") == SYSTEM_OUTPUT_SCHEMA, "Wrong system-output schema")
    require(is_synthetic(system) == is_synthetic(frozen), "System/frozen synthetic marker mismatch")
    require(system.get("label_bundle_sha256") == frozen["label_bundle_sha256"], "System output was not bound to this label freeze")
    raw_records = system.get("records")
    require(isinstance(raw_records, list) and raw_records, "System outputs have no records")
    frozen_cases = frozen.get("cases")
    require(isinstance(frozen_cases, list) and frozen_cases, "Frozen label bundle has no cases")
    labels = {string(case.get("case_id")): case for case in frozen_cases if isinstance(case, dict)}
    require(len(labels) == len(frozen_cases) and "" not in labels, "Frozen case IDs empty/duplicate")
    score_protocol = frozen.get("score_protocol")
    require(isinstance(score_protocol, dict), "Frozen bundle lacks score_protocol")
    protocol_models = unique_strings(score_protocol.get("models"), "frozen score models")
    protocol_methods = unique_strings(score_protocol.get("methods"), "frozen score methods")
    protocol_strata = unique_strings(score_protocol.get("strata"), "frozen score strata")
    require(protocol_models and protocol_methods, "Frozen score axes are empty")
    require(set(protocol_strata) == set(FORMAL_STRATUM_COUNTS), "Frozen score strata are incomplete")
    require(score_protocol.get("require_complete_grid") is True, "Frozen score protocol must require a complete grid")
    bridge_protocol = score_protocol.get("bridge_aware_cleanliness")
    require(
        isinstance(bridge_protocol, dict)
        and bridge_protocol.get("enabled") is True
        and bridge_protocol.get("allowed_node_rule")
        == "frozen repair packet union canonical injection bridge nodes",
        "Frozen bridge-aware cleanliness protocol is missing or changed",
    )
    for supplied, frozen_axis, label in (
        (expected_models, protocol_models, "models"),
        (expected_methods, protocol_methods, "methods"),
        (expected_strata, protocol_strata, "strata"),
    ):
        if supplied is not None:
            require(list(supplied) == frozen_axis, f"Runtime expected {label} differ from frozen score protocol")
    for case_id, case in labels.items():
        require(string(case.get("source_scc_uid")), f"Frozen {case_id} has no source_scc_uid")
        stratum = string(case.get("stratum"))
        require(stratum in FORMAL_STRATUM_COUNTS, f"Frozen {case_id} has invalid stratum")
        require(bool(string(case.get("case_uid"))), f"Frozen {case_id} has no case_uid")
        post_fingerprint = case.get("post_content_fingerprint")
        if stratum.startswith("post_"):
            require_sha256(post_fingerprint, f"frozen/{case_id}.post_content_fingerprint")
        else:
            require(post_fingerprint is None, f"Frozen pre case {case_id} must have null post_content_fingerprint")
        endpoints = unique_strings(case.get("endpoint_node_ids", []), f"frozen/{case_id}.endpoint_node_ids")
        bridges = unique_strings(case.get("bridge_node_ids", []), f"frozen/{case_id}.bridge_node_ids")
        bridge_allowed = unique_strings(
            case.get("bridge_aware_allowed_node_ids", []),
            f"frozen/{case_id}.bridge_aware_allowed_node_ids",
        )
        require(
            set(endpoints) | set(bridges) | set(bridge_allowed)
            <= set(case.get("node_universe") or []),
            f"Frozen {case_id} has endpoint/bridge nodes outside its SCC universe",
        )
        require(
            set(bridge_allowed)
            == set(case.get("repair_packet_node_ids") or []) | set(bridges),
            f"Frozen {case_id} bridge-aware allowed set is inconsistent",
        )
        require(
            case.get("endpoint_metric_available") is bool(endpoints),
            f"Frozen {case_id} endpoint availability flag is inconsistent",
        )
    actual_endpoint_scope = [
        stratum for stratum in FORMAL_STRATUM_COUNTS
        if any(case["stratum"] == stratum for case in frozen_cases)
        and all(
            case["endpoint_metric_available"]
            for case in frozen_cases if case["stratum"] == stratum
        )
    ]
    require(
        score_protocol.get("endpoint_metric_scope") == actual_endpoint_scope,
        "Frozen endpoint scope differs from actual endpoint availability",
    )
    aliases: dict[str, str] = {}
    for case_id, case in labels.items():
        for alias in {
            case_id,
            string(case.get("original_case_id")),
            string(case.get("case_uid")),
            string(case.get("legacy_case_id")),
        } - {""}:
            require(alias not in aliases or aliases[alias] == case_id, f"Ambiguous frozen case alias {alias!r}")
            aliases[alias] = case_id
    seen: set[tuple[str, str, str]] = set()
    records = []
    for raw in raw_records:
        require(isinstance(raw, dict), "System-output record is not an object")
        input_case_id = string(raw.get("case_id"))
        case_id = aliases.get(input_case_id, "")
        model = string(raw.get("model"))
        method = string(raw.get("method"))
        require(case_id in labels and model and method, f"System-output identity fields invalid for {input_case_id!r}")
        key = (case_id, model, method)
        require(key not in seen, f"Duplicate system cell {key}")
        seen.add(key)
        label = labels[case_id]
        require("case_uid" in raw, f"{key}: exact case_uid is required")
        require(
            string(raw.get("case_uid")) == label["case_uid"],
            f"{key}: case_uid mismatch",
        )
        require("post_content_fingerprint" in raw, f"{key}: post_content_fingerprint field is required")
        require(
            raw.get("post_content_fingerprint") == label["post_content_fingerprint"],
            f"{key}: post_content_fingerprint mismatch",
        )
        raw_source_scc_uid = string(raw.get("source_scc_uid"))
        raw_stratum = string(raw.get("stratum"))
        require(raw_source_scc_uid, f"{key}: source_scc_uid is required")
        require(raw_stratum, f"{key}: stratum is required")
        require(raw_source_scc_uid == label["source_scc_uid"], f"{key}: SCC mismatch")
        require(raw_stratum == label["stratum"], f"{key}: stratum mismatch")
        valid = raw.get("output_valid")
        require(type(valid) is bool, f"{key}: output_valid must be boolean")
        prediction = string(raw.get("prediction"))
        require(prediction in PREDICTIONS, f"{key}: prediction invalid")
        require((valid and prediction != "invalid") or (not valid and prediction == "invalid"), f"{key}: validity/prediction inconsistent")
        returned = unique_strings(raw.get("returned_node_ids", []), f"{key}.returned_node_ids")
        require(set(returned) <= set(label["node_universe"]), f"{key}: returned nodes outside frozen SCC identity")
        if not valid:
            require(not returned, f"{key}: invalid output must have an empty returned set")
        returned_set = set(returned)
        evidence = set(label["minimal_evidence_node_ids"])
        repair = set(label["repair_packet_node_ids"])
        endpoint_set = set(label["endpoint_node_ids"])
        bridge_allowed = set(label["bridge_aware_allowed_node_ids"])
        target = label["final_target_conflict"]
        positive = target == "present"
        negative = target == "absent"
        cleanliness_eligible = positive and label["final_evidence_sufficient"] in {"yes", "partial"}
        if cleanliness_eligible:
            require(evidence and repair, f"{key}: eligible positive lacks frozen evidence/repair denominator")
        row = {
            "case_id": case_id,
            "input_case_id": input_case_id,
            "case_uid": label["case_uid"],
            "stratum": label["stratum"],
            "domain": label["domain"],
            "source_scc_uid": label["source_scc_uid"],
            "model": model,
            "method": method,
            "output_valid": valid,
            "prediction": prediction,
            "returned_node_ids": sorted(returned),
            "strict_valid": float(valid),
            "abstained": float(valid and prediction in {"insufficient", "abstain"}),
            "retained_node_count": float(len(returned_set)),
            "positive_detection_hit": float(valid and prediction == "present") if positive else None,
            "negative_safe_hit": float(valid and prediction in {"absent", "insufficient", "abstain"}) if negative else None,
            "over_reported": float(valid and prediction == "present") if negative else None,
            "positive_cleanliness_eligible": float(cleanliness_eligible) if positive else None,
            "expert_evidence_recall": (len(returned_set & evidence) / len(evidence)) if cleanliness_eligible else None,
            "repair_packet_coverage": (len(returned_set & repair) / len(repair)) if cleanliness_eligible else None,
            "excess_node_rate": ((len(returned_set - repair) / len(returned_set)) if returned_set else 0.0) if cleanliness_eligible else None,
            "bridge_aware_excess_node_rate": (
                (len(returned_set - bridge_allowed) / len(returned_set)) if returned_set else 0.0
            ) if cleanliness_eligible else None,
            "bridge_aware_precision": (
                (len(returned_set & bridge_allowed) / len(returned_set)) if returned_set else 0.0
            ) if cleanliness_eligible else None,
            "expert_evidence_density": ((len(returned_set & evidence) / len(returned_set)) if returned_set else 0.0) if cleanliness_eligible else None,
            "endpoint_recall": (len(returned_set & endpoint_set) / len(endpoint_set)) if endpoint_set else None,
            "endpoint_density": ((len(returned_set & endpoint_set) / len(returned_set)) if returned_set else 0.0) if endpoint_set else None,
        }
        records.append(row)

    if score_protocol["require_complete_grid"]:
        models = protocol_models
        methods = protocol_methods
        strata = protocol_strata
        expected_case_ids = {case_id for case_id, case in labels.items() if case["stratum"] in strata}
        require(expected_case_ids, "Expected strata select no frozen cases")
        expected = {
            (case_id, model, method)
            for case_id in expected_case_ids
            for model in models
            for method in methods
        }
        require(seen == expected, f"System grid incomplete: missing={len(expected - seen)}, extra={len(seen - expected)}")

    metric_names = [
        "strict_valid",
        "abstained",
        "retained_node_count",
        "positive_detection_hit",
        "negative_safe_hit",
        "over_reported",
        "positive_cleanliness_eligible",
        "expert_evidence_recall",
        "repair_packet_coverage",
        "excess_node_rate",
        "bridge_aware_excess_node_rate",
        "bridge_aware_precision",
        "expert_evidence_density",
        "endpoint_recall",
        "endpoint_density",
    ]
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        grouped[(row["stratum"], row["model"], row["method"])].append(row)
    summaries = []
    for (stratum, model, method), rows in sorted(grouped.items()):
        require(stratum in FORMAL_STRATUM_COUNTS, f"Unknown output stratum {stratum}")
        metrics = {}
        for metric in metric_names:
            if not any(row.get(metric) is not None for row in rows):
                continue
            group_seed = _stable_group_seed(seed, (stratum, model, method, metric))
            point, low, high, n_rows, n_clusters = cluster_bootstrap_ci(
                rows, metric, samples=bootstrap_samples, seed=group_seed
            )
            metrics[metric] = {
                "mean": point,
                "clustered_ci_95_low": low,
                "clustered_ci_95_high": high,
                "denominator_case_presentations": n_rows,
                "source_scc_clusters": n_clusters,
            }
        summaries.append({
            "stratum": stratum,
            "model": model,
            "method": method,
            "n_system_cells": len(rows),
            "metrics": metrics,
        })
    cleanliness = seal({
        "schema_version": CLEANLINESS_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": is_synthetic(frozen),
        "label_bundle_sha256": frozen["label_bundle_sha256"],
        "system_outputs_file_sha256": file_sha256(system_outputs_path),
        "metric_definitions": {
            "expert_evidence_recall": "|returned ∩ frozen minimal evidence| / |frozen minimal evidence|",
            "repair_packet_coverage": "|returned ∩ frozen repair packet| / |frozen repair packet|",
            "excess_node_rate": "|returned minus frozen repair packet| / |returned|; 0 by convention for an empty returned set on an eligible positive",
            "bridge_aware_excess_node_rate": "|returned minus (frozen repair packet union canonical bridge nodes)| / |returned|; 0 for an empty returned set",
            "bridge_aware_precision": "|returned ∩ (frozen repair packet union canonical bridge nodes)| / |returned|; 0 for an empty returned set",
            "endpoint_density": "|returned ∩ frozen canonical endpoints| / |returned|; reported only for strata with endpoint availability frozen true for every case",
            "cleanliness_eligibility": "positive with final evidence_sufficient in {yes, partial}; every eligible case must have non-empty frozen evidence and repair denominators",
            "invalid_policy": "invalid is retained; recall/coverage/hit are zero on eligible positive cases",
        },
        "records": records,
    })
    summary = seal({
        "schema_version": CLUSTER_SUMMARY_SCHEMA,
        "tool_version": TOOL_VERSION,
        "synthetic_fixture": is_synthetic(frozen),
        "cleanliness_artifact_sha256": cleanliness["artifact_sha256"],
        "cluster_unit": "source_scc_uid",
        "bootstrap_samples": bootstrap_samples,
        "base_seed": seed,
        "cluster_weighting": "equal weight per source_scc_uid after within-SCC metric averaging",
        "strata_never_pooled": list(FORMAL_STRATUM_COUNTS),
        "endpoint_metric_scope": actual_endpoint_scope,
        "complete_grid_contract": {
            "required": True,
            "runtime_flag_requested": require_complete_grid,
            "expected_models": protocol_models,
            "expected_methods": protocol_methods,
            "expected_strata": protocol_strata,
            "source": "frozen_label_bundle.score_protocol",
        },
        "three_experts_are_not_resampled": True,
        "groups": summaries,
    })
    return cleanliness, summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    layout = subparsers.add_parser("layout", help="Audit public manifest/private codebook shape without labels")
    layout.add_argument("--manifest", type=Path, required=True)
    layout.add_argument("--codebook", type=Path, required=True)
    layout.add_argument("--translations", type=Path, required=True)
    layout.add_argument("--out", type=Path, required=True)

    merge = subparsers.add_parser("merge", help="Validate and merge exactly three blind expert exports")
    merge.add_argument("--manifest", type=Path, required=True)
    merge.add_argument("--codebook", type=Path, required=True)
    merge.add_argument("--translations", type=Path, required=True)
    merge.add_argument("--exports", type=Path, nargs=3, required=True, metavar=("EXPERT1", "EXPERT2", "EXPERT3"))
    merge.add_argument("--out-dir", type=Path, required=True)

    freeze = subparsers.add_parser("freeze", help="Freeze labels after required outcome-blind adjudication")
    freeze.add_argument("--merged", type=Path, required=True)
    freeze.add_argument("--adjudication", type=Path, required=True)
    freeze.add_argument("--manifest", type=Path, required=True)
    freeze.add_argument("--codebook", type=Path, required=True)
    freeze.add_argument("--translations", type=Path, required=True)
    freeze.add_argument("--legacy-labels", type=Path, required=True)
    freeze.add_argument("--exports", type=Path, nargs=3, required=True, metavar=("EXPERT1", "EXPERT2", "EXPERT3"))
    freeze.add_argument("--protocol-id", required=True)
    freeze.add_argument("--frozen-by", required=True)
    freeze.add_argument("--frozen-at", required=True)
    freeze.add_argument("--confirm-outcome-blind-freeze", action="store_true")
    freeze.add_argument("--score-models", nargs="+", required=True)
    freeze.add_argument("--score-methods", nargs="+", required=True)
    freeze.add_argument("--score-strata", nargs="+", choices=tuple(FORMAL_STRATUM_COUNTS), required=True)
    freeze.add_argument("--out", type=Path, required=True)

    score = subparsers.add_parser("score", help="Compute frozen-label cleanliness and source-SCC clustered CI")
    score.add_argument("--labels", type=Path, required=True)
    score.add_argument("--system-outputs", type=Path, required=True)
    score.add_argument("--bootstrap-samples", type=int, default=20_000)
    score.add_argument("--seed", type=int, default=20260909)
    score.add_argument("--require-complete-grid", action="store_true")
    score.add_argument("--expected-models", nargs="+")
    score.add_argument("--expected-methods", nargs="+")
    score.add_argument("--expected-strata", nargs="+", choices=tuple(FORMAL_STRATUM_COUNTS))
    score.add_argument("--out-dir", type=Path, required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "layout":
            artifact = layout_audit(args.manifest, args.codebook, args.translations)
            write_json(args.out, artifact)
            print(json.dumps({"ok": True, "out": str(args.out), "counts": artifact["counts"]}, ensure_ascii=False))
        elif args.command == "merge":
            paths = merge_to_dir(
                args.manifest,
                args.codebook,
                args.translations,
                args.exports,
                args.out_dir,
            )
            merged = read_json(paths["merged"])
            adjudication = read_json(paths["adjudication"])
            print(json.dumps({
                "ok": True,
                "out_dir": str(args.out_dir),
                "counts": merged["counts"],
                "adjudication_cases": len(adjudication["cases"]),
            }, ensure_ascii=False))
        elif args.command == "freeze":
            artifact = freeze_labels(
                args.merged,
                args.adjudication,
                args.manifest,
                args.codebook,
                args.translations,
                args.legacy_labels,
                export_paths=args.exports,
                protocol_id=args.protocol_id,
                frozen_by=args.frozen_by,
                frozen_at=args.frozen_at,
                confirm_outcome_blind_freeze=args.confirm_outcome_blind_freeze,
                score_models=args.score_models,
                score_methods=args.score_methods,
                score_strata=args.score_strata,
            )
            write_json(args.out, artifact)
            print(json.dumps({
                "ok": True,
                "out": str(args.out),
                "counts": artifact["counts"],
                "label_bundle_sha256": artifact["label_bundle_sha256"],
            }, ensure_ascii=False))
        elif args.command == "score":
            cleanliness, summary = score_system_outputs(
                args.labels,
                args.system_outputs,
                bootstrap_samples=args.bootstrap_samples,
                seed=args.seed,
                require_complete_grid=args.require_complete_grid,
                expected_models=args.expected_models,
                expected_methods=args.expected_methods,
                expected_strata=args.expected_strata,
            )
            cleanliness_path = args.out_dir / "cleanliness_records.json"
            summary_path = args.out_dir / "clustered_summary.json"
            write_json(cleanliness_path, cleanliness)
            write_json(summary_path, summary)
            print(json.dumps({
                "ok": True,
                "out_dir": str(args.out_dir),
                "records": len(cleanliness["records"]),
                "groups": len(summary["groups"]),
            }, ensure_ascii=False))
        else:  # pragma: no cover - argparse prevents this
            raise AssertionError(args.command)
    except ValidationError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
